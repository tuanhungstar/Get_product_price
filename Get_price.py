import sys
import time
import os
import json
import io
import zipfile
import platform
import subprocess
import re
import base64
import mimetypes
from typing import Optional, List, Dict, Any

import requests
from PyQt6.QtCore import Qt, QThread, pyqtSignal, QProcess
from PyQt6.QtGui import QPixmap
from PyQt6.QtWidgets import (
    QApplication,
    QMainWindow,
    QWidget,
    QVBoxLayout,
    QHBoxLayout,
    QGridLayout,
    QLabel,
    QLineEdit,
    QPushButton,
    QComboBox,
    QTextEdit,
    QFileDialog,
    QGroupBox,
    QSplitter,
    QTabWidget,
    QTableWidget,
    QTableWidgetItem,
    QHeaderView,
    QCheckBox,
    QDialog,
    QMessageBox,
    QSpinBox,
    QDoubleSpinBox,
    QRadioButton,
)

from selenium import webdriver
from selenium.common.exceptions import WebDriverException
from selenium.webdriver.chrome.service import Service as ChromeService
from selenium.webdriver.remote.webelement import WebElement
from selenium.webdriver.common.by import By

# --- JAVASCRIPT LIBRARIES ---

JS_LIST_CONTEXTS = """
const results = { iframes: [], shadow_hosts: [] };
function getXPath(el) {
    let path = [];
    while (el && el.nodeType === 1) {
        let index = 0, sibling = el.previousSibling;
        while (sibling) {
            if (sibling.nodeType === 1 && sibling.tagName === el.tagName) index++;
            sibling = sibling.previousSibling;
        }
        const tagName = (el.tagName || '').toLowerCase();
        const pathIndex = (index > 0) ? `[${index + 1}]` : '';
        path.unshift(tagName + pathIndex);
        el = el.parentNode;
    }
    return path.join('/');
}
function findElements(contextNode, docPath) {
    const iframes = contextNode.querySelectorAll('iframe');
    for (const iframe of iframes) {
        const framePath = getXPath(iframe);
        const fullPath = docPath ? `${docPath} -> IFRAME` : `IFRAME`;
        let info = { xpath: framePath, full_path: `${fullPath}(${framePath})`, id: iframe.id || '', name: iframe.name || '' };
        results.iframes.push(info);
        try {
            if (iframe.contentDocument) findElements(iframe.contentDocument, fullPath);
        } catch (e) { info.error = 'Cross-origin iframe'; }
    }
    const allElements = contextNode.querySelectorAll('*');
    for (const el of allElements) {
        if (el.shadowRoot) {
            const hostPath = getXPath(el);
            const fullPath = docPath ? `${docPath} -> SHADOW_HOST` : `SHADOW_HOST`;
            results.shadow_hosts.push({ host_xpath: hostPath, full_path: `${fullPath}(${hostPath})`, host_id: el.id || '', host_tag: el.tagName.toLowerCase(), mode: el.shadowRoot.mode });
            findElements(el.shadowRoot, fullPath);
        }
    }
}
findElements(document, 'DOCUMENT');
return results;
"""

JS_SMART_SEARCH_ENGINE = """
const searchText = (arguments[0] || '').toLowerCase();
const mode = arguments[1] || 'smart';
const searchContext = arguments.length > 2 ? arguments[2] : document;
if (!searchText) return [];
function findElements(contextNode) {
    let matches = [];
    const checkNode = (el) => {
        try {
            if (el.nodeType !== 1) return false;
            const style = window.getComputedStyle(el);
            if (style.display === 'none' || style.visibility === 'hidden' || style.opacity === '0') return false;
            const text = (el.innerText || el.textContent || el.value || el.placeholder || '').trim().toLowerCase();
            if (mode === 'exact') return text === searchText;
            return text.includes(searchText);
        } catch (e) { return false; }
    };
    const nodes = contextNode.querySelectorAll('*');
    for (const node of nodes) {
        if (checkNode(node)) matches.push(node);
        if (node.shadowRoot) matches = matches.concat(findElements(node.shadowRoot));
    }
    // Don't search iframes if we are already in a specific context
    if (contextNode === document) {
        const iframes = contextNode.querySelectorAll('iframe');
        for (const iframe of iframes) {
            try {
                if (iframe.contentDocument) matches = matches.concat(findElements(iframe.contentDocument));
            } catch (e) {}
        }
    }
    return matches;
}
let allMatches = findElements(searchContext);
if (mode === 'smart' && allMatches.length > 0) {
    const exactMatches = allMatches.filter(el => (el.innerText || el.textContent || el.value || '').trim().toLowerCase() === searchText);
    if (exactMatches.length > 0) allMatches = exactMatches;
}
return allMatches.filter(el => !allMatches.some(other => el !== other && el.contains(other)));
"""

JS_HIGHLIGHT_ELEMENT = "arguments[0].style.outline = `3px solid ${arguments[1] || 'red'}`;"
JS_CLEAR_HIGHLIGHT = "if(arguments[0]) arguments[0].style.outline = '';"
JS_GET_XPATH = "let e=arguments[0],p;for(p=[];e&&e.nodeType==1;e=e.parentNode){let i=1,s=e.previousSibling;for(;s;s=s.previousSibling)1==s.nodeType&&s.tagName==e.tagName&&i++;p.unshift(e.tagName.toLowerCase()+(i>1?`[${i}]`:''))}return p.join('/')"


# --- EXCEL LIST_WEB.XLSX MANAGERS & EDITOR DIALOG ---

def load_list_web_excel(excel_path: str = "list_web.xlsx"):
    """
    Reads list_web.xlsx and returns:
    - items: list of dicts [{'web': str, 'keywords': list[str], 'image_key': str}]
    - priority_domains: list of domain/name strings (Column A)
    """
    abs_path = os.path.abspath(excel_path)
    items = []
    priority_domains = []
    
    if os.path.exists(abs_path):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(abs_path, data_only=True)
            sheet = wb.active
            rows = list(sheet.iter_rows(values_only=True))
            if rows:
                first_cell = str(rows[0][0] or '').strip().lower()
                start_row = 1 if len(rows) > 1 and any(h in first_cell for h in ['web', 'website', 'domain', 'name']) else 0
                for r in rows[start_row:]:
                    if not r or r[0] is None or not str(r[0]).strip():
                        continue
                    web_name = str(r[0]).strip()
                    kw_raw = str(r[1]).strip() if len(r) > 1 and r[1] is not None else ""
                    img_raw = str(r[2]).strip() if len(r) > 2 and r[2] is not None else ""
                    
                    keywords = [k.strip() for k in kw_raw.split(",") if k.strip()] if kw_raw and kw_raw.lower() != "nan" else []
                    image_key = img_raw if img_raw and img_raw.lower() != "nan" else ""
                    
                    priority_domains.append(web_name)
                    items.append({
                        "web": web_name,
                        "keywords": keywords,
                        "image_key": image_key
                    })
        except Exception as e:
            print(f"[WARN] Failed to read {excel_path} via openpyxl: {e}")
            
    # Fallback if list_web.xlsx is missing or returns empty list
    if not priority_domains:
        priority_file = os.path.abspath("priority_web.txt")
        if os.path.exists(priority_file):
            with open(priority_file, "r", encoding="utf-8") as f:
                priority_domains = [line.strip().lower() for line in f if line.strip() and not line.strip().startswith("#")]
                
    return items, priority_domains


def save_list_web_excel(items: list, excel_path: str = "list_web.xlsx") -> tuple:
    """
    Saves the website configuration items back to list_web.xlsx using openpyxl.
    Returns (success_bool, message_str).
    """
    abs_path = os.path.abspath(excel_path)
    temp_path = abs_path + ".tmp"
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Web List"
        ws.append(["web name", "pop up smart search key word", "pop up click image"])
        for item in items:
            web = item.get("web", "")
            kws = item.get("keywords", [])
            kw_str = ", ".join(kws) if isinstance(kws, list) else str(kws)
            img = item.get("image_key", "")
            ws.append([web, kw_str, img])
        
        try:
            wb.save(abs_path)
            return True, "Successfully saved"
        except (PermissionError, BlockingIOError, OSError) as e_lock:
            try:
                wb.save(temp_path)
                if os.path.exists(abs_path):
                    try:
                        os.remove(abs_path)
                    except Exception:
                        pass
                os.replace(temp_path, abs_path)
                return True, "Successfully saved"
            except Exception:
                return False, f"File is locked: {e_lock}"
    except Exception as e:
        print(f"[ERROR] Failed to save {excel_path}: {e}")
        return False, str(e)


class WebListEditorDialog(QDialog):
    def __init__(self, parent=None, excel_path="list_web.xlsx"):
        super().__init__(parent)
        self.excel_path = os.path.abspath(excel_path)
        self.setWindowTitle("🌐 Website & Popup Configuration Editor")
        self.resize(780, 500)
        self._init_ui()
        self._load_data()

    def _init_ui(self):
        layout = QVBoxLayout(self)

        info_label = QLabel(
            "<b>Manage Priority Websites & Popup Dismissal Rules (list_web.xlsx)</b><br/>"
            "• <b>Column A (Web Name / Domain)</b>: Website keyword or domain to prioritize in Google Search.<br/>"
            "• <b>Column B (Smart Search Keywords)</b>: Comma-separated popup button text keywords (e.g. <i>Từ chối, Bữa khác nha, Đóng</i>).<br/>"
            "• <b>Column C (Click Image)</b>: Image key or filename in <code>Click_image/</code> to find and left-click on screen."
        )
        info_label.setWordWrap(True)
        layout.addWidget(info_label)

        self.table = QTableWidget()
        self.table.setColumnCount(3)
        self.table.setHorizontalHeaderLabels(["Web Name / Domain (Col A)", "Popup Smart Keywords (Col B)", "Click Image File / Key (Col C)"])
        self.table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Interactive)
        self.table.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        self.table.horizontalHeader().setSectionResizeMode(2, QHeaderView.ResizeMode.Interactive)
        self.table.setColumnWidth(0, 200)
        self.table.setColumnWidth(2, 200)
        layout.addWidget(self.table)

        btn_layout = QHBoxLayout()
        self.btn_add = QPushButton("➕ Add Row")
        self.btn_add.clicked.connect(self._action_add_row)
        btn_layout.addWidget(self.btn_add)

        self.btn_delete = QPushButton("🗑️ Delete Selected")
        self.btn_delete.clicked.connect(self._action_delete_row)
        btn_layout.addWidget(self.btn_delete)

        self.btn_browse_img = QPushButton("📂 Select Click Image...")
        self.btn_browse_img.clicked.connect(self._action_browse_image)
        btn_layout.addWidget(self.btn_browse_img)

        btn_layout.addStretch()

        self.btn_save = QPushButton("💾 Save to list_web.xlsx")
        self.btn_save.setStyleSheet("font-weight: bold; background-color: #2e7d32; color: white;")
        self.btn_save.clicked.connect(self._action_save)
        btn_layout.addWidget(self.btn_save)

        self.btn_cancel = QPushButton("❌ Cancel")
        self.btn_cancel.clicked.connect(self.reject)
        btn_layout.addWidget(self.btn_cancel)

        layout.addLayout(btn_layout)

    def _load_data(self):
        items, _ = load_list_web_excel(self.excel_path)
        self.table.setRowCount(0)
        for row_idx, item in enumerate(items):
            self.table.insertRow(row_idx)
            self.table.setItem(row_idx, 0, QTableWidgetItem(item.get("web", "")))
            kws = item.get("keywords", [])
            kw_str = ", ".join(kws) if isinstance(kws, list) else str(kws)
            self.table.setItem(row_idx, 1, QTableWidgetItem(kw_str))
            self.table.setItem(row_idx, 2, QTableWidgetItem(item.get("image_key", "")))

    def _action_add_row(self):
        row_idx = self.table.rowCount()
        self.table.insertRow(row_idx)
        self.table.setItem(row_idx, 0, QTableWidgetItem("new_domain"))
        self.table.setItem(row_idx, 1, QTableWidgetItem("Từ chối, Bỏ qua"))
        self.table.setItem(row_idx, 2, QTableWidgetItem(""))

    def _action_delete_row(self):
        curr_row = self.table.currentRow()
        if curr_row >= 0:
            self.table.removeRow(curr_row)

    def _action_browse_image(self):
        curr_row = self.table.currentRow()
        if curr_row < 0:
            QMessageBox.warning(self, "Warning", "Please select a row first in the table.")
            return
        click_img_dir = os.path.abspath("Click_image")
        os.makedirs(click_img_dir, exist_ok=True)
        file_path, _ = QFileDialog.getOpenFileName(
            self, "Select Click Image or Key File", click_img_dir, "Supported Files (*.png *.jpg *.jpeg *.txt);;All Files (*)"
        )
        if file_path:
            rel_name = os.path.basename(file_path)
            if rel_name.endswith('.txt'):
                rel_name = rel_name[:-4]
            self.table.setItem(curr_row, 2, QTableWidgetItem(rel_name))

    def _action_save(self):
        items = []
        for r in range(self.table.rowCount()):
            web_item = self.table.item(r, 0)
            kw_item = self.table.item(r, 1)
            img_item = self.table.item(r, 2)

            web = web_item.text().strip() if web_item else ""
            if not web:
                continue
            kw_str = kw_item.text().strip() if kw_item else ""
            keywords = [k.strip() for k in kw_str.split(",") if k.strip()]
            image_key = img_item.text().strip() if img_item else ""

            items.append({
                "web": web,
                "keywords": keywords,
                "image_key": image_key
            })

        success, msg = save_list_web_excel(items, self.excel_path)
        if success:
            QMessageBox.information(self, "Success", f"Successfully saved {len(items)} website rule(s) to list_web.xlsx!")
            self.accept()
        else:
            QMessageBox.critical(
                self,
                "Save Error",
                f"Could not save list_web.xlsx:\n\n{msg}\n\n"
                "Please close list_web.xlsx if it is open in Microsoft Excel or another editor and try saving again."
            )


# --- CHROMEDRIVER AUTOMATIC DOWNLOAD UTILITIES ---

def safe_requests_get(url: str, timeout: int = 15, **kwargs) -> requests.Response:
    """Helper to perform GET requests with SSL verification fallback."""
    try:
        return requests.get(url, timeout=timeout, **kwargs)
    except requests.exceptions.SSLError:
        return requests.get(url, timeout=timeout, verify=False, **kwargs)


def get_installed_chrome_version() -> Optional[str]:
    """Detects the installed version of Google Chrome across Windows, macOS, and Linux."""
    system = platform.system()

    if system == "Windows":
        try:
            import winreg
            keys = [
                (winreg.HKEY_CURRENT_USER, r"Software\Google\Chrome\BLBeacon"),
                (winreg.HKEY_LOCAL_MACHINE, r"SOFTWARE\Google\Chrome\BLBeacon"),
                (winreg.HKEY_LOCAL_MACHINE, r"SOFTWARE\WOW6432Node\Google\Chrome\BLBeacon"),
            ]
            for root, subkey in keys:
                try:
                    with winreg.OpenKey(root, subkey) as key:
                        val, _ = winreg.QueryValueEx(key, "version")
                        if val:
                            return str(val).strip()
                except Exception:
                    pass
        except ImportError:
            pass

        ps_paths = [
            r"C:\Program Files\Google\Chrome\Application\chrome.exe",
            r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
            os.path.expandvars(r"%LOCALAPPDATA%\Google\Chrome\Application\chrome.exe"),
        ]
        for p in ps_paths:
            if os.path.exists(p):
                try:
                    cmd = f'powershell -NoProfile -Command "(Get-Item \'{p}\').VersionInfo.ProductVersion"'
                    out = subprocess.check_output(cmd, shell=True, text=True, stderr=subprocess.DEVNULL)
                    match = re.search(r'[\d.]+', out)
                    if match:
                        return match.group(0)
                except Exception:
                    pass

    elif system == "Darwin":
        mac_paths = [
            "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
            "/Applications/Chromium.app/Contents/MacOS/Chromium",
        ]
        for p in mac_paths:
            if os.path.exists(p):
                try:
                    out = subprocess.check_output([p, "--version"], text=True, stderr=subprocess.DEVNULL)
                    match = re.search(r'[\d.]+', out)
                    if match:
                        return match.group(0)
                except Exception:
                    pass
        plist_path = "/Applications/Google Chrome.app/Contents/Info.plist"
        if os.path.exists(plist_path):
            try:
                import plistlib
                with open(plist_path, "rb") as fp:
                    pl = plistlib.load(fp)
                    ver = pl.get("CFBundleShortVersionString")
                    if ver:
                        return str(ver).strip()
            except Exception:
                pass

    elif system == "Linux":
        commands = [
            ["google-chrome", "--version"],
            ["google-chrome-stable", "--version"],
            ["chromium", "--version"],
            ["chromium-browser", "--version"]
        ]
        for cmd in commands:
            try:
                out = subprocess.check_output(cmd, text=True, stderr=subprocess.DEVNULL)
                match = re.search(r'[\d.]+', out)
                if match:
                    return match.group(0)
            except Exception:
                pass

    return None


def get_platform_key() -> str:
    """Returns the Chrome for Testing platform identifier."""
    sys_name = platform.system()
    machine = platform.machine().lower()

    if sys_name == "Windows":
        is_64 = sys.maxsize > 2**32 or "64" in machine
        return "win64" if is_64 else "win32"
    elif sys_name == "Darwin":
        return "mac-arm64" if "arm" in machine or "aarch" in machine else "mac-x64"
    elif sys_name == "Linux":
        return "linux-arm64" if "arm" in machine or "aarch" in machine else "linux64"
    return "win64"


def find_chromedriver_url(chrome_version: str) -> tuple[Optional[str], Optional[str]]:
    """Resolves ChromeDriver download URL and version string for a given Chrome version."""
    major_ver = chrome_version.split('.')[0]
    platform_key = get_platform_key()

    if int(major_ver) >= 115:
        # 1. Per-milestone API lookup
        try:
            url = 'https://googlechromelabs.github.io/chrome-for-testing/latest-versions-per-milestone-with-downloads.json'
            r = safe_requests_get(url, timeout=10)
            if r.status_code == 200:
                data = r.json()
                milestone = data.get('milestones', {}).get(major_ver)
                if milestone:
                    downloads = milestone.get('downloads', {}).get('chromedriver', [])
                    for item in downloads:
                        if item.get('platform') == platform_key:
                            return item.get('url'), milestone.get('version')
        except Exception:
            pass

        # 2. LATEST_RELEASE_<MAJOR> lookup
        try:
            lr_url = f'https://googlechromelabs.github.io/chrome-for-testing/LATEST_RELEASE_{major_ver}'
            r = safe_requests_get(lr_url, timeout=10)
            if r.status_code == 200:
                version = r.text.strip()
                dl_url = f'https://storage.googleapis.com/chrome-for-testing-public/{version}/{platform_key}/chromedriver-{platform_key}.zip'
                return dl_url, version
        except Exception:
            pass

        # 3. Last known good Stable channel fallback
        try:
            url = 'https://googlechromelabs.github.io/chrome-for-testing/last-known-good-versions-with-downloads.json'
            r = safe_requests_get(url, timeout=10)
            if r.status_code == 200:
                data = r.json()
                stable = data.get('channels', {}).get('Stable', {})
                version = stable.get('version')
                downloads = stable.get('downloads', {}).get('chromedriver', [])
                for item in downloads:
                    if item.get('platform') == platform_key:
                        return item.get('url'), version
        except Exception:
            pass

    else: # Chrome version < 115
        try:
            lr_url = f'https://chromedriver.storage.googleapis.com/LATEST_RELEASE_{major_ver}'
            r = safe_requests_get(lr_url, timeout=10)
            if r.status_code == 200:
                version = r.text.strip()
                leg_plat = 'win32' if platform.system() == 'Windows' else ('mac64' if platform.system() == 'Darwin' else 'linux64')
                dl_url = f'https://chromedriver.storage.googleapis.com/{version}/chromedriver_{leg_plat}.zip'
                return dl_url, version
        except Exception:
            pass

    return None, None


class ChromeDriverDownloaderThread(QThread):
    log_signal = pyqtSignal(str, str) # message, level
    finished_signal = pyqtSignal(bool, str) # success, path_or_error

    def run(self):
        try:
            self.log_signal.emit("🔍 Detecting installed Google Chrome version...", "INFO")
            chrome_version = get_installed_chrome_version()

            if chrome_version:
                self.log_signal.emit(f"Detected Google Chrome version: {chrome_version}", "SUCCESS")
                dl_url, version_str = find_chromedriver_url(chrome_version)
            else:
                self.log_signal.emit("Could not detect local Google Chrome version automatically. Attempting latest stable fallback...", "WARN")
                dl_url, version_str = find_chromedriver_url("120.0.0.0")

            if not dl_url:
                self.finished_signal.emit(False, "Failed to resolve a compatible ChromeDriver download URL.")
                return

            self.log_signal.emit(f"Found ChromeDriver v{version_str}. Target platform: {get_platform_key()}", "INFO")
            self.log_signal.emit(f"Downloading: {dl_url}", "INFO")

            resp = safe_requests_get(dl_url, timeout=30)
            if resp.status_code != 200:
                self.finished_signal.emit(False, f"Download failed with HTTP status code {resp.status_code}")
                return

            self.log_signal.emit("Download complete. Extracting ChromeDriver binary...", "INFO")
            z = zipfile.ZipFile(io.BytesIO(resp.content))
            exe_name = "chromedriver.exe" if platform.system() == "Windows" else "chromedriver"

            found_member = None
            for member in z.namelist():
                if os.path.basename(member).lower() == exe_name.lower():
                    found_member = member
                    break

            if not found_member:
                self.finished_signal.emit(False, f"Could not find '{exe_name}' inside downloaded zip.")
                return

            target_path = os.path.abspath(exe_name)
            with z.open(found_member) as source, open(target_path, "wb") as target:
                target.write(source.read())

            if platform.system() != "Windows":
                try:
                    os.chmod(target_path, 0o755)
                except Exception:
                    pass
                if platform.system() == "Darwin":
                    try:
                        subprocess.run(["xattr", "-d", "com.apple.quarantine", target_path], stderr=subprocess.DEVNULL)
                    except Exception:
                        pass

            self.finished_signal.emit(True, target_path)

        except Exception as e:
            self.finished_signal.emit(False, str(e))


class GoogleSearchThread(QThread):
    log_signal = pyqtSignal(str, str) # message, level
    results_signal = pyqtSignal(bool, list, str) # success, items_list, error_msg

    def __init__(self, api_key: str, search_engine_id: str, query: str):
        super().__init__()
        self.api_key = api_key
        self.search_engine_id = search_engine_id
        self.query = query

    def run(self):
        if not self.api_key:
            self.results_signal.emit(False, [], "Google Search API Key is missing. Please configure it in the Configuration tab.")
            return
        if not self.search_engine_id:
            self.results_signal.emit(False, [], "Search Engine ID (CX) is missing. Please configure it in the Configuration tab.")
            return
        if not self.query:
            self.results_signal.emit(False, [], "Search query cannot be empty.")
            return

        self.log_signal.emit(f"Executing Google Custom Search for: '{self.query}'...", "INFO")
        url = "https://www.googleapis.com/customsearch/v1"
        params = {
            "key": self.api_key,
            "cx": self.search_engine_id,
            "q": self.query,
            "num": 10
        }

        try:
            resp = safe_requests_get(url, params=params, timeout=15)
            if resp.status_code != 200:
                err_data = resp.json().get("error", {}) if resp.headers.get("content-type", "").startswith("application/json") else {}
                err_msg = err_data.get("message", f"HTTP {resp.status_code}: {resp.text[:200]}")
                self.results_signal.emit(False, [], f"Google Search API Error: {err_msg}")
                return

            data = resp.json()
            items = data.get("items", [])
            parsed_results = []
            for item in items:
                parsed_results.append({
                    "title": item.get("title", ""),
                    "snippet": item.get("snippet", "").replace("\n", " "),
                    "link": item.get("link", ""),
                })

            self.results_signal.emit(True, parsed_results, "")
        except Exception as e:
            self.results_signal.emit(False, [], str(e))


class GeminiApiThread(QThread):
    log_signal = pyqtSignal(str, str) # message, level
    response_signal = pyqtSignal(bool, str) # success, response_text_or_error

    def __init__(self, api_key: str, model: str, prompt: str, media_path: Optional[str] = None, retry_delay: float = 5.0, max_retries: int = 3):
        super().__init__()
        self.api_key = api_key
        self.model = model or "gemini-2.5-flash"
        self.prompt = prompt
        self.media_path = media_path
        self.retry_delay = retry_delay
        self.max_retries = max_retries

    def run(self):
        if not self.api_key:
            self.response_signal.emit(False, "Gemini API Key is missing. Please configure it in the Configuration tab.")
            return
        if not self.prompt and not self.media_path:
            self.response_signal.emit(False, "Prompt text or media file must be provided.")
            return

        self.log_signal.emit(f"Calling Gemini API (Model: {self.model})...", "INFO")

        url = f"https://generativelanguage.googleapis.com/v1beta/models/{self.model}:generateContent?key={self.api_key}"
        parts = []

        if self.prompt:
            parts.append({"text": self.prompt})

        if self.media_path and os.path.exists(self.media_path):
            try:
                mime_type, _ = mimetypes.guess_type(self.media_path)
                if not mime_type:
                    ext = os.path.splitext(self.media_path)[1].lower()
                    mime_type = {
                        ".pdf": "application/pdf",
                        ".png": "image/png",
                        ".jpg": "image/jpeg",
                        ".jpeg": "image/jpeg",
                        ".webp": "image/webp",
                    }.get(ext, "application/octet-stream")

                self.log_signal.emit(f"Attaching file: {os.path.basename(self.media_path)} (MIME: {mime_type})...", "INFO")
                with open(self.media_path, "rb") as f:
                    encoded_b64 = base64.b64encode(f.read()).decode("utf-8")

                parts.append({
                    "inline_data": {
                        "mime_type": mime_type,
                        "data": encoded_b64
                    }
                })
            except Exception as e:
                self.response_signal.emit(False, f"Failed to read media file: {e}")
                return

        payload = {
            "contents": [
                {
                    "parts": parts
                }
            ]
        }

        headers = {"Content-Type": "application/json"}
        for attempt in range(1, self.max_retries + 2):
            try:
                resp = requests.post(url, json=payload, headers=headers, timeout=60)
                if resp.status_code == 200:
                    res_json = resp.json()
                    candidates = res_json.get("candidates", [])
                    if not candidates:
                        self.response_signal.emit(False, "Gemini API returned no candidates.")
                        return

                    text_response = ""
                    for candidate in candidates:
                        content = candidate.get("content", {})
                        for part in content.get("parts", []):
                            if "text" in part:
                                text_response += part["text"]

                    if not text_response:
                        text_response = "[No text output returned by model]"

                    self.response_signal.emit(True, text_response)
                    return
                else:
                    err_info = resp.json().get("error", {}) if "json" in resp.headers.get("content-type", "") else {}
                    err_msg = err_info.get("message", f"HTTP {resp.status_code}: {resp.text[:300]}")
                    if attempt <= self.max_retries:
                        self.log_signal.emit(f"Gemini API Error/Busy ({err_msg}). Retrying in {self.retry_delay}s (Attempt {attempt}/{self.max_retries})...", "WARN")
                        time.sleep(self.retry_delay)
                    else:
                        self.response_signal.emit(False, f"Gemini API Error: {err_msg}")
                        return
            except Exception as e:
                if attempt <= self.max_retries:
                    self.log_signal.emit(f"Gemini API Exception ({e}). Retrying in {self.retry_delay}s (Attempt {attempt}/{self.max_retries})...", "WARN")
                    time.sleep(self.retry_delay)
                else:
                    self.response_signal.emit(False, f"Gemini API Exception: {e}")
                    return


def parse_numeric_price(price_str: str) -> float:
    if not price_str or not isinstance(price_str, str):
        return float('inf')
    lower = price_str.lower().strip()
    if lower in ["not_found", "n/a", "none", "-", "stopped", "failed / not found", "", "error"]:
        return float('inf')
    digits_only = re.sub(r'[^\d]', '', price_str)
    if digits_only:
        try:
            return float(digits_only)
        except ValueError:
            pass
    return float('inf')


class GetPriceWorkflowThread(QThread):
    log_signal = pyqtSignal(str, str) # text, level
    progress_signal = pyqtSignal(str) # status text
    search_table_signal = pyqtSignal(list) # items for search table update
    finished_signal = pyqtSignal(bool, list, str, str) # success, results_list, error_msg, final_prompt

    def __init__(self, product_name: str, config: dict, process_first_only: bool = True):
        super().__init__()
        self.product_name = product_name
        self.config = config
        self.process_first_only = process_first_only

    def run(self):
        final_prompt = ""
        results_list = []
        try:
            search_key = self.config.get("google_search_api_key", "").strip()
            cx = self.config.get("google_search_engine_id", "").strip()
            gemini_key = self.config.get("gemini_api_key", "").strip()
            gemini_model = self.config.get("gemini_model", "gemini-2.5-flash").strip()
            driver_path = self.config.get("chromedriver_path", "").strip()
            delay_between_calls = float(self.config.get("gemini_delay_between_calls", 2.0))
            retry_delay = float(self.config.get("gemini_retry_delay", 5.0))
            max_retries = int(self.config.get("gemini_max_retries", 3))

            if not search_key or not cx:
                err = "Google Search API Key or CX is missing. Please configure them in the Configuration tab."
                self.progress_signal.emit("⚠️ Search API Key or CX Missing")
                self.finished_signal.emit(False, [], err, "")
                return

            if not gemini_key:
                err = "Gemini API Key is missing. Please configure it in the Configuration tab."
                self.progress_signal.emit("⚠️ Gemini API Key Missing")
                self.finished_signal.emit(False, [], err, "")
                return

            # Auto-detect driver path if not explicitly provided
            if not driver_path or not os.path.exists(driver_path):
                default_driver = os.path.abspath("chromedriver.exe" if platform.system() == "Windows" else "chromedriver")
                if os.path.exists(default_driver):
                    driver_path = default_driver

            if not driver_path or not os.path.exists(driver_path):
                err = "ChromeDriver executable not found. Please set or download ChromeDriver in the Configuration tab."
                self.progress_signal.emit("⚠️ ChromeDriver Missing")
                self.finished_signal.emit(False, [], err, "")
                return

            # --- STEP 1: Google Custom Search ---
            self.progress_signal.emit(f"🔍 Searching Google for: '{self.product_name}'...")
            self.log_signal.emit(f"Step 1: Searching Google for '{self.product_name}'...", "HEADING")

            search_url = "https://www.googleapis.com/customsearch/v1"
            params = {"key": search_key, "cx": cx, "q": self.product_name, "num": 10}
            resp = safe_requests_get(search_url, params=params, timeout=15)

            if resp.status_code != 200:
                err_data = resp.json().get("error", {}) if "json" in resp.headers.get("content-type", "") else {}
                err_msg = err_data.get("message", f"HTTP {resp.status_code}: {resp.text[:200]}")
                self.progress_signal.emit("❌ Google Search API Error")
                self.log_signal.emit(f"Google Search failed: {err_msg}", "ERROR")
                self.finished_signal.emit(False, [], f"Google Search failed: {err_msg}", "")
                return

            items = resp.json().get("items", [])
            if not items:
                self.progress_signal.emit("❌ No Google Search Results Found")
                self.log_signal.emit(f"No Google Search results found for '{self.product_name}'. Stopping workflow.", "WARN")
                self.finished_signal.emit(False, [], f"No search results found on Google for '{self.product_name}'.", "")
                return

            self.log_signal.emit(f"Google Search returned {len(items)} result(s).", "SUCCESS")

            # --- STEP 2: Filter by list_web.xlsx / priority_web.txt ---
            self.progress_signal.emit("🌐 Filtering search results by priority websites...")
            self.log_signal.emit("Step 2: Filtering results via list_web.xlsx / priority_web.txt...", "INFO")

            excel_configs, priority_keywords = load_list_web_excel("list_web.xlsx")
            self.log_signal.emit(f"Loaded {len(priority_keywords)} priority domain keyword(s).", "INFO")

            filtered_items = []
            for item in items:
                link = item.get("link", "")
                display_link = item.get("displayLink", "")
                title = item.get("title", "")
                snippet = item.get("snippet", "").replace("\n", " ")

                link_lower = link.lower()
                disp_lower = display_link.lower()
                matches_priority = any(k.lower() in link_lower or k.lower() in disp_lower for k in priority_keywords) if priority_keywords else True
                if matches_priority:
                    domain_name = display_link or (link.split("/")[2] if "://" in link else link)
                    filtered_items.append({
                        "title": title,
                        "snippet": snippet,
                        "link": link,
                        "domain": domain_name
                    })

            # Emit search items to update Search Test tab table
            table_items = filtered_items if filtered_items else [{
                "title": item.get("title", ""),
                "snippet": item.get("snippet", "").replace("\n", " "),
                "link": item.get("link", "")
            } for item in items]

            self.search_table_signal.emit(table_items)

            if not filtered_items:
                msg = "No search results matched priority websites listed in list_web.xlsx."
                self.progress_signal.emit("⚠️ No Priority Website Matched")
                self.log_signal.emit(msg, "WARN")
                self.finished_signal.emit(False, [], msg, "")
                return

            target_items = filtered_items[:1] if self.process_first_only else filtered_items
            self.log_signal.emit(f"Selected {len(target_items)} priority link(s) to process.", "SUCCESS")

            # Prepare Prompt & Popup Texts
            prompt_file = os.path.abspath("prompt.txt")
            raw_prompt_template = ""
            if os.path.exists(prompt_file):
                with open(prompt_file, "r", encoding="utf-8") as f:
                    raw_prompt_template = f.read()
            else:
                raw_prompt_template = "Target Product: 'product_search'\nExtract product_name, price, status, similarity as JSON."

            final_prompt = raw_prompt_template.replace("product_search", self.product_name)

            close_popup_file = os.path.abspath("close_popup.txt")
            close_texts = []
            if os.path.exists(close_popup_file):
                with open(close_popup_file, "r", encoding="utf-8") as f:
                    close_texts = [line.strip() for line in f if line.strip() and not line.strip().startswith("#")]

            screenshot_dir = os.path.abspath("screenshot")
            os.makedirs(screenshot_dir, exist_ok=True)

            total_targets = len(target_items)
            for idx, item in enumerate(target_items, start=1):
                target_link = item["link"]
                item_domain = item["domain"]
                self.progress_signal.emit(f"🌐 [{idx}/{total_targets}] Opening: {item_domain}...")
                self.log_signal.emit(f"Step 3 ({idx}/{total_targets}): Opening website {target_link}", "INFO")

                service = ChromeService(executable_path=driver_path)
                options = webdriver.ChromeOptions()
                options.add_argument("--start-maximized")
                driver = None
                screenshot_path = ""

                try:
                    driver = webdriver.Chrome(service=service, options=options)
                    driver.get(target_link)
                    time.sleep(3)

                    # Look up popup config for the current website domain
                    link_lower = target_link.lower()
                    domain_lower = item_domain.lower()
                    matched_config = None
                    for cfg in excel_configs:
                        w_pattern = cfg.get("web", "").strip().lower()
                        if w_pattern and (w_pattern in link_lower or w_pattern in domain_lower):
                            matched_config = cfg
                            break

                    site_keywords = matched_config["keywords"] if (matched_config and matched_config.get("keywords")) else close_texts
                    
                    if site_keywords:
                        self.progress_signal.emit(f"🧹 [{idx}/{total_targets}] Dismissing popups via keywords ({', '.join(site_keywords)})...")
                        for popup_text in site_keywords:
                            try:
                                elements = driver.execute_script(JS_SMART_SEARCH_ENGINE, popup_text, 'smart')
                                if elements:
                                    el = elements[0]
                                    try:
                                        el.click()
                                    except Exception:
                                        driver.execute_script("arguments[0].click();", el)
                                    time.sleep(1)
                            except Exception:
                                pass

                    if matched_config and matched_config.get("image_key"):
                        img_key = matched_config["image_key"]
                        self.progress_signal.emit(f"🖼️ [{idx}/{total_targets}] Dismissing popup via visual image match ('{img_key}')...")
                        try:
                            from BOT_take_image import Old_utility
                            bot_util = Old_utility()
                            res = bot_util.left_click(img_key)
                            self.log_signal.emit(f"Visual popup click result for '{img_key}': {res}", "INFO")
                            time.sleep(1)
                        except Exception as e_img:
                            self.log_signal.emit(f"Visual image click error for '{img_key}': {e_img}", "WARN")

                    self.progress_signal.emit(f"📷 [{idx}/{total_targets}] Capturing screenshot...")
                    timestamp_str = time.strftime("%Y%m%d_%H%M%S")
                    screenshot_path = os.path.join(screenshot_dir, f"price_{idx}_{timestamp_str}.png")
                    driver.save_screenshot(screenshot_path)
                    self.log_signal.emit(f"Saved screenshot: {screenshot_path}", "SUCCESS")

                except Exception as e_driver:
                    self.log_signal.emit(f"Error visiting {target_link}: {e_driver}", "ERROR")
                finally:
                    if driver:
                        try:
                            driver.quit()
                        except Exception:
                            pass

                if not screenshot_path or not os.path.exists(screenshot_path):
                    continue

                # Inter-call delay before calling Gemini API (when processing subsequent items)
                if idx > 1 and delay_between_calls > 0:
                    self.progress_signal.emit(f"⏳ Waiting {delay_between_calls}s delay before Gemini API call ({idx}/{total_targets})...")
                    self.log_signal.emit(f"Inter-call delay: Sleeping {delay_between_calls}s before calling Gemini API...", "INFO")
                    time.sleep(delay_between_calls)

                # Call Gemini API with automatic retry handling
                self.progress_signal.emit(f"🤖 [{idx}/{total_targets}] Analyzing with Gemini AI...")
                self.log_signal.emit(f"Step 6 ({idx}/{total_targets}): Sending screenshot to Gemini API...", "INFO")

                try:
                    with open(screenshot_path, "rb") as sf:
                        encoded_b64 = base64.b64encode(sf.read()).decode("utf-8")

                    gemini_url = f"https://generativelanguage.googleapis.com/v1beta/models/{gemini_model}:generateContent?key={gemini_key}"
                    payload = {
                        "contents": [
                            {
                                "parts": [
                                    {"text": final_prompt},
                                    {
                                        "inline_data": {
                                            "mime_type": "image/png",
                                            "data": encoded_b64
                                        }
                                    }
                                ]
                            }
                        ]
                    }

                    headers = {"Content-Type": "application/json"}
                    raw_text_out = ""
                    parsed_dict = {}

                    for attempt in range(1, max_retries + 2):
                        try:
                            self.log_signal.emit(f"Gemini API request attempt {attempt}/{max_retries + 1}...", "INFO")
                            gemini_resp = requests.post(gemini_url, json=payload, headers=headers, timeout=60)

                            if gemini_resp.status_code == 200:
                                res_json = gemini_resp.json()
                                for candidate in res_json.get("candidates", []):
                                    for part in candidate.get("content", {}).get("parts", []):
                                        if "text" in part:
                                            raw_text_out += part["text"]

                                try:
                                    match = re.search(r'\{.*\}', raw_text_out, re.DOTALL)
                                    if match:
                                        parsed_dict = json.loads(match.group(0))
                                    else:
                                        parsed_dict = json.loads(raw_text_out)
                                except Exception:
                                    parsed_dict = {"product_name": "Parsing error", "price": "N/A", "status": "N/A", "similarity": "0%"}
                                break
                            else:
                                err_info = gemini_resp.json().get("error", {}) if "json" in gemini_resp.headers.get("content-type", "") else {}
                                err_msg = err_info.get("message", f"HTTP {gemini_resp.status_code}: {gemini_resp.text[:200]}")
                                raw_text_out = f"Gemini API Error: {err_msg}"
                                if attempt <= max_retries:
                                    self.progress_signal.emit(f"⚠️ Gemini API Busy ({gemini_resp.status_code}). Retrying in {retry_delay}s ({attempt}/{max_retries})...")
                                    self.log_signal.emit(f"Gemini API returned status {gemini_resp.status_code} ({err_msg}). Retrying in {retry_delay}s...", "WARN")
                                    time.sleep(retry_delay)
                                else:
                                    parsed_dict = {"product_name": f"API Error ({gemini_resp.status_code})", "price": "N/A", "status": "Error", "similarity": "0%"}
                        except Exception as e_req:
                            raw_text_out = f"Gemini API Exception: {e_req}"
                            if attempt <= max_retries:
                                self.progress_signal.emit(f"⚠️ Gemini Exception. Retrying in {retry_delay}s ({attempt}/{max_retries})...")
                                self.log_signal.emit(f"Gemini API exception ({e_req}). Retrying in {retry_delay}s...", "WARN")
                                time.sleep(retry_delay)
                            else:
                                parsed_dict = {"product_name": "Connection Error", "price": "N/A", "status": "Error", "similarity": "0%"}

                    num_price = parse_numeric_price(parsed_dict.get("price", ""))
                    results_list.append({
                        "index": idx,
                        "link": target_link,
                        "title": item.get("title", ""),
                        "domain": item_domain,
                        "product_name": parsed_dict.get("product_name", "-"),
                        "price": parsed_dict.get("price", "-"),
                        "status": parsed_dict.get("status", "-"),
                        "similarity": parsed_dict.get("similarity", "-"),
                        "raw_text": raw_text_out,
                        "screenshot_path": screenshot_path,
                        "numeric_price": num_price,
                        "final_prompt": final_prompt
                    })

                    self.log_signal.emit(f"Extracted price for [{item_domain}]: {parsed_dict.get('price')}", "SUCCESS")

                except Exception as e_gem:
                    self.log_signal.emit(f"Gemini API request error for {target_link}: {e_gem}", "ERROR")

            if results_list:
                self.progress_signal.emit("✅ Price extraction complete!")
                self.finished_signal.emit(True, results_list, "", final_prompt)
            else:
                self.progress_signal.emit("❌ Failed to extract prices from links.")
                self.finished_signal.emit(False, [], "Failed to process links.", final_prompt)

        except Exception as e:
            self.progress_signal.emit(f"❌ Error: {e}")
            self.finished_signal.emit(False, [], str(e), final_prompt)


REQUIRED_EXCEL_COLUMNS = [
    'ID',
    'Project name',
    'Product Name',
    'Cheapest price',
    'Web link of cheapest',
    'screenshot of cheapest',
    'Expensive price',
    'Web link of expensive',
    'screenshot of expensive',
    'Json all',
    'status'
]


def validate_and_read_excel(excel_path: str):
    """
    Validates if excel_path matches required Get_price_list.xlsx template schema.
    Returns: (is_valid: bool, error_message: str, headers: list[str], rows_data: list[dict])
    """
    if not os.path.exists(excel_path):
        return False, f"File does not exist: {excel_path}", [], []

    try:
        import openpyxl
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        sheet = wb.active
        all_rows = list(sheet.iter_rows(values_only=True))

        if not all_rows:
            return False, "Excel file is empty.", [], []

        header_row = [str(cell).strip() if cell is not None else "" for cell in all_rows[0]]

        # Check required columns
        missing_cols = []
        for req in REQUIRED_EXCEL_COLUMNS:
            if not any(req.lower() == h.lower() for h in header_row):
                missing_cols.append(req)

        if missing_cols:
            err_msg = (
                "Invalid Excel Template!\n\n"
                "Missing required column(s):\n• " + "\n• ".join(missing_cols) + "\n\n"
                f"Expected Template Columns:\n{', '.join(REQUIRED_EXCEL_COLUMNS)}"
            )
            return False, err_msg, header_row, []

        rows_data = []
        for row_idx, row_values in enumerate(all_rows[1:], start=2): # openpyxl row index starts at 2 for data
            if not row_values or all(cell is None or str(cell).strip() == "" for cell in row_values):
                continue

            row_dict = {}
            for col_i, header_name in enumerate(header_row):
                if header_name:
                    val = row_values[col_i] if col_i < len(row_values) else ""
                    row_dict[header_name] = str(val if val is not None else "").strip()

            rows_data.append({
                "excel_row_num": row_idx,
                "ui_index": len(rows_data),
                "data": row_dict
            })

        return True, "", header_row, rows_data

    except Exception as e:
        return False, f"Error reading Excel file: {e}", [], []


class GetPriceAllWorkerThread(QThread):
    log_signal = pyqtSignal(str, str) # text, level
    progress_signal = pyqtSignal(str) # status text
    row_started_signal = pyqtSignal(int) # ui row index
    row_finished_signal = pyqtSignal(int, dict) # ui row index, updated row dict
    finished_signal = pyqtSignal(bool, str) # success, message

    def __init__(self, excel_path: str, rows_to_process: list, config: dict):
        super().__init__()
        self.excel_path = excel_path
        self.rows_to_process = rows_to_process
        self.config = config
        self._is_paused = False
        self._is_stopped = False

    def pause(self):
        self._is_paused = True

    def resume(self):
        self._is_paused = False

    def stop(self):
        self._is_stopped = True
        self._is_paused = False

    def run(self):
        import openpyxl
        try:
            search_key = self.config.get("google_search_api_key", "").strip()
            cx = self.config.get("google_search_engine_id", "").strip()
            gemini_key = self.config.get("gemini_api_key", "").strip()
            gemini_model = self.config.get("gemini_model", "gemini-2.5-flash").strip()
            driver_path = self.config.get("chromedriver_path", "").strip()
            delay_between_calls = float(self.config.get("gemini_delay_between_calls", 2.0))
            retry_delay = float(self.config.get("gemini_retry_delay", 5.0))
            max_retries = int(self.config.get("gemini_max_retries", 3))

            if not search_key or not cx:
                self.finished_signal.emit(False, "Google Search API Key or CX is missing. Please configure them in Configuration tab.")
                return

            if not gemini_key:
                self.finished_signal.emit(False, "Gemini API Key is missing. Please configure it in Configuration tab.")
                return

            if not driver_path or not os.path.exists(driver_path):
                default_driver = os.path.abspath("chromedriver.exe" if platform.system() == "Windows" else "chromedriver")
                if os.path.exists(default_driver):
                    driver_path = default_driver

            if not driver_path or not os.path.exists(driver_path):
                self.finished_signal.emit(False, "ChromeDriver executable not found. Please set ChromeDriver in Configuration tab.")
                return

            excel_configs, priority_keywords = load_list_web_excel("list_web.xlsx")
            prompt_file = os.path.abspath("prompt.txt")
            raw_prompt_template = ""
            if os.path.exists(prompt_file):
                with open(prompt_file, "r", encoding="utf-8") as f:
                    raw_prompt_template = f.read()
            else:
                raw_prompt_template = "Target Product: 'product_search'\nExtract product_name, price, status, similarity as JSON."

            close_popup_file = os.path.abspath("close_popup.txt")
            close_texts = []
            if os.path.exists(close_popup_file):
                with open(close_popup_file, "r", encoding="utf-8") as f:
                    close_texts = [line.strip() for line in f if line.strip() and not line.strip().startswith("#")]

            screenshot_dir = os.path.abspath("screenshot")
            os.makedirs(screenshot_dir, exist_ok=True)

            total_rows = len(self.rows_to_process)
            for idx, item_info in enumerate(self.rows_to_process, start=1):
                if self._is_stopped:
                    self.log_signal.emit("Batch processing stopped by user.", "WARN")
                    break

                while self._is_paused:
                    if self._is_stopped:
                        break
                    time.sleep(0.5)

                if self._is_stopped:
                    break

                excel_row_num = item_info["excel_row_num"]
                row_data = item_info["data"]
                row_index_ui = item_info["ui_index"]

                # Check status column: skip if status == "done"
                current_status = str(row_data.get("status", "")).strip().lower()
                if current_status == "done":
                    self.log_signal.emit(f"Row {excel_row_num} status is 'done'. Skipping.", "INFO")
                    continue

                product_name = str(row_data.get("Product Name", "")).strip()
                row_id = str(row_data.get("ID", "")).strip()
                project_name = str(row_data.get("Project name", "")).strip()

                if not product_name or product_name.lower() == "nan":
                    self.log_signal.emit(f"Row {excel_row_num} has no Product Name. Skipping.", "WARN")
                    continue

                self.row_started_signal.emit(row_index_ui)
                self.progress_signal.emit(f"[{idx}/{total_rows}] Processing Row {excel_row_num}: '{product_name}'...")
                self.log_signal.emit(f"Step 1 ({idx}/{total_rows}): Processing Row {excel_row_num} (ID: {row_id}, Project: {project_name}) for '{product_name}'", "HEADING")

                # Step A: Search Google Custom Search
                search_url = "https://www.googleapis.com/customsearch/v1"
                params = {"key": search_key, "cx": cx, "q": product_name, "num": 10}
                resp = safe_requests_get(search_url, params=params, timeout=15)

                if resp.status_code != 200:
                    self.log_signal.emit(f"Google Search failed for row {excel_row_num}: HTTP {resp.status_code}", "ERROR")
                    continue

                search_items = resp.json().get("items", [])
                if not search_items:
                    self.log_signal.emit(f"No Google Search results found for '{product_name}'", "WARN")
                    continue

                # Step B: Filter items by priority domains
                filtered_items = []
                for s_item in search_items:
                    link = s_item.get("link", "")
                    display_link = s_item.get("displayLink", "")
                    title = s_item.get("title", "")
                    snippet = s_item.get("snippet", "").replace("\n", " ")

                    link_lower = link.lower()
                    disp_lower = display_link.lower()
                    matched_kw = None
                    if priority_keywords:
                        for k in priority_keywords:
                            if k.lower() in link_lower or k.lower() in disp_lower:
                                matched_kw = k
                                break
                    else:
                        matched_kw = display_link or "web"

                    if matched_kw or not priority_keywords:
                        domain_name = display_link or (link.split("/")[2] if "://" in link else link)
                        filtered_items.append({
                            "title": title,
                            "snippet": snippet,
                            "link": link,
                            "domain": domain_name,
                            "matched_kw": matched_kw or domain_name
                        })

                if not filtered_items:
                    self.log_signal.emit(f"No search results matched priority filter for '{product_name}'", "WARN")
                    continue

                self.log_signal.emit(f"Row {excel_row_num}: Found {len(filtered_items)} matching priority link(s).", "SUCCESS")

                final_prompt = raw_prompt_template.replace("product_search", product_name)
                row_link_results = []

                # Step C: Process ALL priority links for this row
                for link_idx, f_item in enumerate(filtered_items, start=1):
                    if self._is_stopped:
                        break

                    while self._is_paused:
                        if self._is_stopped:
                            break
                        time.sleep(0.5)

                    if self._is_stopped:
                        break

                    target_link = f_item["link"]
                    item_domain = f_item["domain"]
                    web_priority_name = f_item["matched_kw"]

                    # Screenshot naming format: ID_project_web
                    def sanitize(val: str) -> str:
                        return re.sub(r'[\\/*?:"<>|]', '_', str(val)).strip()

                    clean_id = sanitize(row_id) or f"row{excel_row_num}"
                    clean_proj = sanitize(project_name) or "proj"
                    clean_web = sanitize(web_priority_name) or "web"

                    if len(filtered_items) > 1:
                        ss_filename = f"{clean_id}_{clean_proj}_{clean_web}_{link_idx}.png"
                    else:
                        ss_filename = f"{clean_id}_{clean_proj}_{clean_web}.png"

                    screenshot_path = os.path.join(screenshot_dir, ss_filename)

                    self.progress_signal.emit(f"🌐 Row {excel_row_num} [{link_idx}/{len(filtered_items)}]: Opening {web_priority_name}...")
                    self.log_signal.emit(f"Opening link ({link_idx}/{len(filtered_items)}): {target_link}", "INFO")

                    service = ChromeService(executable_path=driver_path)
                    options = webdriver.ChromeOptions()
                    options.add_argument("--start-maximized")
                    driver = None

                    try:
                        driver = webdriver.Chrome(service=service, options=options)
                        driver.get(target_link)
                        time.sleep(3)

                        link_lower = target_link.lower()
                        domain_lower = item_domain.lower()
                        matched_config = None
                        for cfg in excel_configs:
                            w_pattern = cfg.get("web", "").strip().lower()
                            if w_pattern and (w_pattern in link_lower or w_pattern in domain_lower):
                                matched_config = cfg
                                break

                        site_keywords = matched_config["keywords"] if (matched_config and matched_config.get("keywords")) else close_texts

                        if site_keywords:
                            for popup_text in site_keywords:
                                try:
                                    elements = driver.execute_script(JS_SMART_SEARCH_ENGINE, popup_text, 'smart')
                                    if elements:
                                        el = elements[0]
                                        try:
                                            el.click()
                                        except Exception:
                                            driver.execute_script("arguments[0].click();", el)
                                        time.sleep(1)
                                except Exception:
                                    pass

                        if matched_config and matched_config.get("image_key"):
                            img_key = matched_config["image_key"]
                            try:
                                from BOT_take_image import Old_utility
                                bot_util = Old_utility()
                                res = bot_util.left_click(img_key)
                                time.sleep(1)
                            except Exception:
                                pass

                        driver.save_screenshot(screenshot_path)
                        self.log_signal.emit(f"Saved screenshot: {ss_filename}", "SUCCESS")

                    except Exception as e_driver:
                        self.log_signal.emit(f"Error visiting {target_link}: {e_driver}", "ERROR")
                    finally:
                        if driver:
                            try:
                                driver.quit()
                            except Exception:
                                pass

                    if not screenshot_path or not os.path.exists(screenshot_path):
                        continue

                    if link_idx > 1 and delay_between_calls > 0:
                        time.sleep(delay_between_calls)

                    # Send screenshot to Gemini API
                    self.progress_signal.emit(f"🤖 Row {excel_row_num} [{link_idx}/{len(filtered_items)}]: Analyzing with Gemini AI...")
                    try:
                        with open(screenshot_path, "rb") as sf:
                            encoded_b64 = base64.b64encode(sf.read()).decode("utf-8")

                        gemini_url = f"https://generativelanguage.googleapis.com/v1beta/models/{gemini_model}:generateContent?key={gemini_key}"
                        payload = {
                            "contents": [
                                {
                                    "parts": [
                                        {"text": final_prompt},
                                        {"inline_data": {"mime_type": "image/png", "data": encoded_b64}}
                                    ]
                                }
                            ]
                        }
                        headers = {"Content-Type": "application/json"}
                        raw_text_out = ""
                        parsed_dict = {}

                        for attempt in range(1, max_retries + 2):
                            try:
                                gemini_resp = requests.post(gemini_url, json=payload, headers=headers, timeout=60)
                                if gemini_resp.status_code == 200:
                                    res_json = gemini_resp.json()
                                    for candidate in res_json.get("candidates", []):
                                        for part in candidate.get("content", {}).get("parts", []):
                                            if "text" in part:
                                                raw_text_out += part["text"]
                                    try:
                                        match = re.search(r'\{.*\}', raw_text_out, re.DOTALL)
                                        if match:
                                            parsed_dict = json.loads(match.group(0))
                                        else:
                                            parsed_dict = json.loads(raw_text_out)
                                    except Exception:
                                        parsed_dict = {"product_name": "Parsing error", "price": "N/A", "status": "N/A", "similarity": "0%"}
                                    break
                                else:
                                    if attempt <= max_retries:
                                        time.sleep(retry_delay)
                                    else:
                                        parsed_dict = {"product_name": f"API Error ({gemini_resp.status_code})", "price": "N/A", "status": "Error", "similarity": "0%"}
                            except Exception as e_req:
                                if attempt <= max_retries:
                                    time.sleep(retry_delay)
                                else:
                                    parsed_dict = {"product_name": "Connection Error", "price": "N/A", "status": "Error", "similarity": "0%"}

                        num_price = parse_numeric_price(parsed_dict.get("price", ""))
                        row_link_results.append({
                            "link": target_link,
                            "domain": item_domain,
                            "web_name": web_priority_name,
                            "product_name": parsed_dict.get("product_name", "-"),
                            "price": parsed_dict.get("price", "-"),
                            "status": parsed_dict.get("status", "-"),
                            "similarity": parsed_dict.get("similarity", "-"),
                            "screenshot_file": ss_filename,
                            "numeric_price": num_price,
                            "raw_gemini": raw_text_out
                        })
                    except Exception as e_gem:
                        self.log_signal.emit(f"Gemini API error for {target_link}: {e_gem}", "ERROR")

                if not row_link_results:
                    self.log_signal.emit(f"Row {excel_row_num}: No extracted price results.", "WARN")
                    continue

                # Step D: Determine cheapest and most expensive items
                valid_items = [r for r in row_link_results if r["numeric_price"] != float('inf')]
                if valid_items:
                    cheapest_item = min(valid_items, key=lambda x: x["numeric_price"])
                    expensive_item = max(valid_items, key=lambda x: x["numeric_price"])
                else:
                    cheapest_item = row_link_results[0]
                    expensive_item = row_link_results[0]

                cheapest_price_str = cheapest_item["price"]
                cheapest_link_str = cheapest_item["link"]
                cheapest_ss_str = cheapest_item["screenshot_file"]

                expensive_price_str = expensive_item["price"]
                expensive_link_str = expensive_item["link"]
                expensive_ss_str = expensive_item["screenshot_file"]

                json_all_str = json.dumps(row_link_results, ensure_ascii=False, indent=2)

                # Update row data dictionary for UI table
                updated_row_dict = dict(row_data)
                updated_row_dict["Cheapest price"] = cheapest_price_str
                updated_row_dict["Web link of cheapest"] = cheapest_link_str
                updated_row_dict["screenshot of cheapest"] = cheapest_ss_str
                updated_row_dict["Expensive price"] = expensive_price_str
                updated_row_dict["Web link of expensive"] = expensive_link_str
                updated_row_dict["screenshot of expensive"] = expensive_ss_str
                updated_row_dict["Json all"] = json_all_str
                updated_row_dict["status"] = "done"

                # Step E: Save to Excel immediately using openpyxl
                try:
                    wb = openpyxl.load_workbook(self.excel_path)
                    sheet = wb.active

                    header_map = {}
                    for col_idx in range(1, sheet.max_column + 1):
                        val = sheet.cell(row=1, column=col_idx).value
                        if val:
                            header_map[str(val).strip().lower()] = col_idx

                    col_updates = {
                        "cheapest price": cheapest_price_str,
                        "web link of cheapest": cheapest_link_str,
                        "screenshot of cheapest": cheapest_ss_str,
                        "expensive price": expensive_price_str,
                        "web link of expensive": expensive_link_str,
                        "screenshot of expensive": expensive_ss_str,
                        "json all": json_all_str,
                        "status": "done"
                    }
                    for col_name_lower, cell_val in col_updates.items():
                        if col_name_lower in header_map:
                            sheet.cell(row=excel_row_num, column=header_map[col_name_lower], value=cell_val)

                    wb.save(self.excel_path)
                    self.log_signal.emit(f"Row {excel_row_num} updated & saved to Excel: {os.path.basename(self.excel_path)}", "SUCCESS")
                except Exception as e_excel:
                    self.log_signal.emit(f"Failed to update Excel for row {excel_row_num}: {e_excel}", "ERROR")

                # Emit GUI row finished signal
                self.row_finished_signal.emit(row_index_ui, updated_row_dict)

            self.finished_signal.emit(True, "Batch processing finished.")
        except Exception as e_main:
            self.finished_signal.emit(False, f"Batch processing error: {e_main}")


CONFIG_FILE_PATH = os.path.abspath("config.json")

DEFAULT_CONFIG = {
    "chromedriver_path": "",
    "working_folder": os.getcwd(),
    "gemini_api_key": "",
    "gemini_model": "gemini-2.5-flash",
    "gemini_delay_between_calls": 2.0,
    "gemini_retry_delay": 5.0,
    "gemini_max_retries": 3,
    "google_search_api_key": "",
    "google_search_engine_id": "",
    "target_url": "https://www.google.com",
}


class ChromeDriverTesterApp(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("ChromeDriver Interactive Tester & Price Extractor v11")
        self.resize(1200, 950)
        self.driver: Optional[webdriver.Chrome] = None
        self.found_elements_cache: List[WebElement] = []
        self.context_cache: List[Dict[str, Any]] = []
        self.highlighted_element: Optional[WebElement] = None
        self.download_thread: Optional[ChromeDriverDownloaderThread] = None
        self.search_thread: Optional[GoogleSearchThread] = None
        self.gemini_thread: Optional[GeminiApiThread] = None
        self.get_price_thread: Optional[GetPriceWorkflowThread] = None

        self.get_price_results: List[Dict[str, Any]] = []
        self.cheapest_idx: int = -1
        self.expensive_idx: int = -1

        # Batch Processing Attributes for Tab 0 (Get Price All)
        self.batch_excel_path: str = ""
        self.batch_rows_data: list = []
        self.batch_headers: list = []
        self.batch_thread: Optional[GetPriceAllWorkerThread] = None

        self._init_ui()

    def _init_ui(self):
        central_widget = QWidget()
        self.setCentralWidget(central_widget)
        main_layout = QVBoxLayout(central_widget)

        self.tab_widget = QTabWidget()

        # ==========================================
        # TAB 0: Get Price All Workflow (Primary Batch Tab)
        # ==========================================
        tab_get_price_all = self._init_tab_get_price_all()
        self.tab_widget.addTab(tab_get_price_all, "📊 Get Price All")

        # ==========================================
        # TAB 1: Get Price Workflow (Single Search Tab)
        # ==========================================
        tab_get_price = QWidget()
        get_price_layout = QVBoxLayout(tab_get_price)

        input_group = QGroupBox("1. Enter Product Search Query & Execution Mode")
        input_layout = QGridLayout(input_group)

        input_layout.addWidget(QLabel("Product Name:"), 0, 0)
        self.get_price_product_input = QLineEdit()
        self.get_price_product_input.setPlaceholderText("Enter exact product name (e.g. iPhone 16 Pro Max 256GB)...")
        self.get_price_product_input.returnPressed.connect(self._action_start_get_price_workflow)
        input_layout.addWidget(self.get_price_product_input, 0, 1)

        self.btn_start_get_price = QPushButton("🚀 Get Price")
        self.btn_start_get_price.clicked.connect(self._action_start_get_price_workflow)
        input_layout.addWidget(self.btn_start_get_price, 0, 2)

        self.chk_process_first_only = QCheckBox("Process 1st Priority Link Only")
        self.chk_process_first_only.setChecked(True)
        self.chk_process_first_only.setToolTip("Checked: process top matching link only. Unchecked: process all matching priority links sequentially.")
        input_layout.addWidget(self.chk_process_first_only, 1, 0, 1, 1)

        self.get_price_status_label = QLabel("Ready to search.")
        self.get_price_status_label.setStyleSheet("font-weight: bold; color: #2b5797;")
        input_layout.addWidget(self.get_price_status_label, 1, 1, 1, 2)

        get_price_layout.addWidget(input_group)

        results_splitter = QSplitter(Qt.Orientation.Horizontal)

        info_group = QGroupBox("2. Extracted Product Information")
        info_layout = QGridLayout(info_group)

        info_layout.addWidget(QLabel("Select Processed Link:"), 0, 0)
        self.combo_processed_links = QComboBox()
        self.combo_processed_links.setToolTip("Select a processed link from dropdown to inspect its extracted details and screenshot")
        self.combo_processed_links.currentIndexChanged.connect(self._on_processed_link_changed)
        info_layout.addWidget(self.combo_processed_links, 0, 1)

        price_bar_layout = QHBoxLayout()
        self.btn_select_cheapest = QPushButton("🏷️ Show Cheapest")
        self.btn_select_cheapest.setToolTip("Jump to link with the lowest extracted price")
        self.btn_select_cheapest.clicked.connect(self._select_cheapest_link)
        price_bar_layout.addWidget(self.btn_select_cheapest)

        self.btn_select_expensive = QPushButton("💎 Show Most Expensive")
        self.btn_select_expensive.setToolTip("Jump to link with the highest extracted price")
        self.btn_select_expensive.clicked.connect(self._select_expensive_link)
        price_bar_layout.addWidget(self.btn_select_expensive)

        self.lbl_price_summary = QLabel("Cheapest: - | Most Expensive: -")
        self.lbl_price_summary.setStyleSheet("font-size: 11px;")
        price_bar_layout.addWidget(self.lbl_price_summary)
        price_bar_layout.addStretch()
        info_layout.addLayout(price_bar_layout, 1, 0, 1, 2)

        info_layout.addWidget(QLabel("Extracted Product Name:"), 2, 0)
        self.lbl_extracted_name = QLabel("-")
        self.lbl_extracted_name.setWordWrap(True)
        self.lbl_extracted_name.setStyleSheet("font-size: 13px; font-weight: bold;")
        info_layout.addWidget(self.lbl_extracted_name, 2, 1)

        info_layout.addWidget(QLabel("Main Price:"), 3, 0)
        self.lbl_extracted_price = QLabel("-")
        self.lbl_extracted_price.setStyleSheet("font-size: 18px; font-weight: bold; color: #d9534f;")
        info_layout.addWidget(self.lbl_extracted_price, 3, 1)

        info_layout.addWidget(QLabel("Availability Status:"), 4, 0)
        self.lbl_extracted_status = QLabel("-")
        self.lbl_extracted_status.setStyleSheet("font-weight: bold; color: #5cb85c;")
        info_layout.addWidget(self.lbl_extracted_status, 4, 1)

        info_layout.addWidget(QLabel("Name Similarity:"), 5, 0)
        self.lbl_extracted_similarity = QLabel("-")
        self.lbl_extracted_similarity.setStyleSheet("font-weight: bold;")
        info_layout.addWidget(self.lbl_extracted_similarity, 5, 1)

        info_layout.addWidget(QLabel("Website Link:"), 6, 0)
        self.lbl_extracted_link = QLineEdit("-")
        self.lbl_extracted_link.setReadOnly(True)
        info_layout.addWidget(self.lbl_extracted_link, 6, 1)

        info_layout.addWidget(QLabel("Raw Gemini Output:"), 7, 0, 1, 2)
        self.txt_raw_gemini_output = QTextEdit()
        self.txt_raw_gemini_output.setReadOnly(True)
        self.txt_raw_gemini_output.setMaximumHeight(140)
        info_layout.addWidget(self.txt_raw_gemini_output, 8, 0, 1, 2)

        results_splitter.addWidget(info_group)

        preview_group = QGroupBox("3. Webpage Screenshot Preview")
        preview_layout = QVBoxLayout(preview_group)

        self.lbl_screenshot_preview = QLabel("No screenshot captured yet.")
        self.lbl_screenshot_preview.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.lbl_screenshot_preview.setStyleSheet("border: 1px dashed gray; background-color: #f8f9fa;")
        preview_layout.addWidget(self.lbl_screenshot_preview)

        btn_open_folder = QPushButton("📂 Open Screenshot Folder")
        btn_open_folder.clicked.connect(self._open_screenshot_folder)
        preview_layout.addWidget(btn_open_folder)

        results_splitter.addWidget(preview_group)
        results_splitter.setSizes([550, 450])

        get_price_layout.addWidget(results_splitter)

        self.tab_widget.addTab(tab_get_price, "🏷️ Get Price")

        # ==========================================
        # TAB 2: Browser & Tester
        # ==========================================
        tab_tester = QWidget()
        tester_layout = QVBoxLayout(tab_tester)

        browser_group = QGroupBox("1. Target URL Setup")
        browser_layout = QGridLayout(browser_group)
        browser_layout.addWidget(QLabel("Target URL:"), 0, 0)
        self.url_input = QLineEdit("https://www.google.com")
        browser_layout.addWidget(self.url_input, 0, 1)

        btn_box = QHBoxLayout()
        self.btn_open = QPushButton("Open Browser / Navigate")
        btn_box.addWidget(self.btn_open)
        self.btn_close = QPushButton("Close Browser")
        btn_box.addWidget(self.btn_close)
        browser_layout.addLayout(btn_box, 0, 2)
        tester_layout.addWidget(browser_group)

        context_group = QGroupBox("2. Set Page Context for Searching")
        context_layout = QGridLayout(context_group)
        self.btn_scan_contexts = QPushButton("📊 Scan for Contexts (Iframes & Shadow DOMs)")
        context_layout.addWidget(self.btn_scan_contexts, 0, 0, 1, 2)
        context_layout.addWidget(QLabel("Search In:"), 1, 0)
        self.context_combo = QComboBox()
        context_layout.addWidget(self.context_combo, 1, 1)
        tester_layout.addWidget(context_group)

        locator_group = QGroupBox("3. Find Elements (within selected context)")
        locator_layout = QGridLayout(locator_group)
        locator_layout.addWidget(QLabel("Locator By:"), 0, 0)
        self.combo_by = QComboBox()
        self.combo_by.addItems(["Smart Search (Text)", "XPATH", "CSS_SELECTOR", "ID", "NAME"])
        locator_layout.addWidget(self.combo_by, 0, 1)
        locator_layout.addWidget(QLabel("Selector / Text:"), 1, 0)
        self.selector_input = QLineEdit()
        locator_layout.addWidget(self.selector_input, 1, 1)
        self.btn_find = QPushButton("🔍 Find Matching Elements")
        locator_layout.addWidget(self.btn_find, 2, 0, 1, 2)
        tester_layout.addWidget(locator_group)

        inspector_group = QGroupBox("4. Inspect & Act on Selected Element")
        inspector_layout = QGridLayout(inspector_group)
        inspector_layout.addWidget(QLabel("Found Elements:"), 0, 0)
        self.results_combo = QComboBox()
        inspector_layout.addWidget(self.results_combo, 0, 1)
        inspector_layout.addWidget(QLabel("Selected's XPath:"), 1, 0)
        self.selected_locator_input = QLineEdit()
        self.selected_locator_input.setReadOnly(True)
        inspector_layout.addWidget(self.selected_locator_input, 1, 1)
        action_btn_layout = QHBoxLayout()

        self.btn_click = QPushButton("👆 Click")
        action_btn_layout.addWidget(self.btn_click)
        self.btn_get_text = QPushButton("📝 Get Text")
        action_btn_layout.addWidget(self.btn_get_text)
        self.btn_send_keys = QPushButton("⌨️ Send Keys")
        action_btn_layout.addWidget(self.btn_send_keys)

        self.value_input = QLineEdit()
        self.value_input.setPlaceholderText("Text for Send Keys")
        action_btn_layout.addWidget(self.value_input)

        inspector_layout.addLayout(action_btn_layout, 2, 0, 1, 2)
        tester_layout.addWidget(inspector_group)

        self.tab_widget.addTab(tab_tester, "🌐 Browser & Tester")

        # ==========================================
        # TAB 3: Configuration
        # ==========================================
        tab_config = QWidget()
        config_layout = QVBoxLayout(tab_config)

        # 1. Environment & Paths Settings
        paths_group = QGroupBox("1. Environment & Path Configuration")
        paths_layout = QGridLayout(paths_group)

        # Working Folder
        paths_layout.addWidget(QLabel("Working Folder:"), 0, 0)
        self.working_folder_input = QLineEdit()
        self.working_folder_input.setPlaceholderText("Select application working directory...")
        paths_layout.addWidget(self.working_folder_input, 0, 1)

        self.btn_browse_working_folder = QPushButton("📁 Browse Folder...")
        self.btn_browse_working_folder.clicked.connect(self._browse_working_folder)
        paths_layout.addWidget(self.btn_browse_working_folder, 0, 2)

        # ChromeDriver Path
        paths_layout.addWidget(QLabel("ChromeDriver Path:"), 1, 0)
        self.driver_path_input = QLineEdit()
        paths_layout.addWidget(self.driver_path_input, 1, 1)

        btn_path_layout = QHBoxLayout()
        self.btn_browse = QPushButton("Browse File...")
        btn_path_layout.addWidget(self.btn_browse)

        self.btn_download_driver = QPushButton("⬇️ Auto-Download ChromeDriver")
        self.btn_download_driver.setToolTip("Detect installed Google Chrome version and download matching ChromeDriver automatically")
        btn_path_layout.addWidget(self.btn_download_driver)

        paths_layout.addLayout(btn_path_layout, 1, 2)
        config_layout.addWidget(paths_group)

        # 2. Google Gemini API Settings
        gemini_group = QGroupBox("2. Google Gemini API Configuration")
        gemini_layout = QGridLayout(gemini_group)
        gemini_layout.addWidget(QLabel("Gemini API Key:"), 0, 0)

        self.gemini_key_input = QLineEdit()
        self.gemini_key_input.setPlaceholderText("Enter Gemini API Key (e.g. AIzaSy...)")
        self.gemini_key_input.setEchoMode(QLineEdit.EchoMode.Password)

        self.btn_toggle_gemini_key = QPushButton("👁️ Show")
        self.btn_toggle_gemini_key.setCheckable(True)
        self.btn_toggle_gemini_key.clicked.connect(self._toggle_gemini_key_visibility)

        gemini_key_box = QHBoxLayout()
        gemini_key_box.addWidget(self.gemini_key_input)
        gemini_key_box.addWidget(self.btn_toggle_gemini_key)
        gemini_layout.addLayout(gemini_key_box, 0, 1)

        gemini_layout.addWidget(QLabel("Gemini Model:"), 1, 0)
        self.gemini_model_combo = QComboBox()
        self.gemini_model_combo.setEditable(True)
        self.gemini_model_combo.addItems([
            "gemini-2.5-flash",
            "gemini-2.5-pro",
            "gemini-2.0-flash",
            "gemini-1.5-flash",
            "gemini-1.5-pro",
            "gemini-1.0-pro"
        ])
        gemini_layout.addWidget(self.gemini_model_combo, 1, 1)

        gemini_layout.addWidget(QLabel("Delay between API Calls (sec):"), 2, 0)
        self.spin_gemini_delay_between_calls = QDoubleSpinBox()
        self.spin_gemini_delay_between_calls.setRange(0.0, 120.0)
        self.spin_gemini_delay_between_calls.setSingleStep(0.5)
        self.spin_gemini_delay_between_calls.setValue(2.0)
        self.spin_gemini_delay_between_calls.setToolTip("Delay time in seconds between two consecutive Gemini API calls")
        gemini_layout.addWidget(self.spin_gemini_delay_between_calls, 2, 1)

        gemini_layout.addWidget(QLabel("Retry Delay on Busy/Error (sec):"), 3, 0)
        self.spin_gemini_retry_delay = QDoubleSpinBox()
        self.spin_gemini_retry_delay.setRange(0.0, 120.0)
        self.spin_gemini_retry_delay.setSingleStep(1.0)
        self.spin_gemini_retry_delay.setValue(5.0)
        self.spin_gemini_retry_delay.setToolTip("Delay time in seconds before retrying when Gemini API responds busy, rate limited, or error")
        gemini_layout.addWidget(self.spin_gemini_retry_delay, 3, 1)

        gemini_layout.addWidget(QLabel("Max Retries on Error:"), 4, 0)
        self.spin_gemini_max_retries = QSpinBox()
        self.spin_gemini_max_retries.setRange(0, 20)
        self.spin_gemini_max_retries.setSingleStep(1)
        self.spin_gemini_max_retries.setValue(3)
        self.spin_gemini_max_retries.setToolTip("Maximum number of retry attempts if Gemini API request fails or returns busy")
        gemini_layout.addWidget(self.spin_gemini_max_retries, 4, 1)

        config_layout.addWidget(gemini_group)

        # 3. Google Custom Search API Settings
        search_group = QGroupBox("3. Google Custom Search API Configuration")
        search_layout = QGridLayout(search_group)
        search_layout.addWidget(QLabel("Google Search API Key:"), 0, 0)
        self.search_key_input = QLineEdit()
        self.search_key_input.setPlaceholderText("Enter Google Custom Search API Key")
        search_layout.addWidget(self.search_key_input, 0, 1)

        search_layout.addWidget(QLabel("Search Engine ID (CX):"), 1, 0)
        self.search_cx_input = QLineEdit()
        self.search_cx_input.setPlaceholderText("Enter Custom Search Engine ID (CX)")
        search_layout.addWidget(self.search_cx_input, 1, 1)
        config_layout.addWidget(search_group)

        # 4. Persistence Actions
        storage_group = QGroupBox("4. Configuration File Actions")
        storage_layout = QHBoxLayout(storage_group)

        self.btn_save_config = QPushButton("💾 Save Config")
        self.btn_save_config.setToolTip("Save current settings to config.json")
        storage_layout.addWidget(self.btn_save_config)

        self.btn_load_config = QPushButton("📂 Load Config File...")
        self.btn_load_config.setToolTip("Load settings from a custom JSON file")
        storage_layout.addWidget(self.btn_load_config)

        self.btn_reset_config = QPushButton("🔄 Load Defaults")
        self.btn_reset_config.setToolTip("Reset all settings to default values")
        storage_layout.addWidget(self.btn_reset_config)

        config_layout.addWidget(storage_group)

        # 5. External Tools & Utilities
        tools_group = QGroupBox("5. External Tools & Utilities")
        tools_layout = QHBoxLayout(tools_group)

        self.btn_take_image = QPushButton("📷 Take image")
        self.btn_take_image.setToolTip("Minimize main application and open BOT_take_image tool")
        tools_layout.addWidget(self.btn_take_image)

        self.btn_manage_web_list = QPushButton("🌐 Priority Websites & Popup Rules")
        self.btn_manage_web_list.setToolTip("Manage priority website domain filters and popup closing rules in list_web.xlsx")
        tools_layout.addWidget(self.btn_manage_web_list)

        config_layout.addWidget(tools_group)
        config_layout.addStretch()

        self.tab_widget.addTab(tab_config, "⚙️ Configuration")

        # ==========================================
        # TAB 4: Google Search Test
        # ==========================================
        tab_search = QWidget()
        search_tab_layout = QVBoxLayout(tab_search)

        search_input_group = QGroupBox("1. Perform Google Custom Search")
        search_input_layout = QGridLayout(search_input_group)

        search_input_layout.addWidget(QLabel("Search Query:"), 0, 0)
        self.search_query_input = QLineEdit()
        self.search_query_input.setPlaceholderText("Enter search query here...")
        self.search_query_input.returnPressed.connect(self._action_run_google_search)
        search_input_layout.addWidget(self.search_query_input, 0, 1)

        self.btn_run_search = QPushButton("🔍 Search Google")
        self.btn_run_search.clicked.connect(self._action_run_google_search)
        search_input_layout.addWidget(self.btn_run_search, 0, 2)
        search_tab_layout.addWidget(search_input_group)

        results_group = QGroupBox("2. Search Results")
        results_layout = QVBoxLayout(results_group)

        self.search_results_table = QTableWidget()
        self.search_results_table.setColumnCount(3)
        self.search_results_table.setHorizontalHeaderLabels(["Title", "Snippet", "Link / URL"])
        self.search_results_table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Interactive)
        self.search_results_table.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        self.search_results_table.horizontalHeader().setSectionResizeMode(2, QHeaderView.ResizeMode.Interactive)
        self.search_results_table.setColumnWidth(0, 250)
        self.search_results_table.setColumnWidth(2, 300)
        self.search_results_table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)

        results_layout.addWidget(self.search_results_table)
        search_tab_layout.addWidget(results_group)

        self.tab_widget.addTab(tab_search, "🔍 Google Search Test")

        # ==========================================
        # TAB 5: Google Gemini Test
        # ==========================================
        tab_gemini = QWidget()
        gemini_tab_layout = QVBoxLayout(tab_gemini)

        prompt_group = QGroupBox("1. Prompt & Optional Media File (Image / PDF)")
        prompt_layout = QGridLayout(prompt_group)

        prompt_layout.addWidget(QLabel("Prompt Text:"), 0, 0)
        self.gemini_prompt_input = QTextEdit()
        self.gemini_prompt_input.setPlaceholderText("Enter prompt instructions or question for Gemini...")
        self.gemini_prompt_input.setMaximumHeight(90)
        prompt_layout.addWidget(self.gemini_prompt_input, 0, 1, 1, 2)

        prompt_layout.addWidget(QLabel("Attach Media File:"), 1, 0)
        self.gemini_media_input = QLineEdit()
        self.gemini_media_input.setPlaceholderText("Path to Image (.png, .jpg, .webp) or Document (.pdf)")
        prompt_layout.addWidget(self.gemini_media_input, 1, 1)

        media_btn_layout = QHBoxLayout()
        self.btn_browse_media = QPushButton("📎 Browse File...")
        self.btn_browse_media.clicked.connect(self._browse_gemini_media_file)
        media_btn_layout.addWidget(self.btn_browse_media)

        self.btn_clear_media = QPushButton("❌ Clear File")
        self.btn_clear_media.clicked.connect(lambda: self.gemini_media_input.clear())
        media_btn_layout.addWidget(self.btn_clear_media)
        prompt_layout.addLayout(media_btn_layout, 1, 2)

        self.btn_send_gemini = QPushButton("🚀 Send Prompt to Gemini")
        self.btn_send_gemini.clicked.connect(self._action_send_gemini_prompt)
        prompt_layout.addWidget(self.btn_send_gemini, 2, 0, 1, 3)

        gemini_tab_layout.addWidget(prompt_group)

        response_group = QGroupBox("2. Gemini Response Output")
        response_layout = QVBoxLayout(response_group)

        self.gemini_output_text = QTextEdit()
        self.gemini_output_text.setReadOnly(True)
        self.gemini_output_text.setPlaceholderText("Gemini response text will appear here...")
        response_layout.addWidget(self.gemini_output_text)

        gemini_tab_layout.addWidget(response_group)

        self.tab_widget.addTab(tab_gemini, "🤖 Google Gemini Test")

        # Splitter: Tabs on top, Activity Log at bottom
        splitter = QSplitter(Qt.Orientation.Vertical)
        splitter.addWidget(self.tab_widget)

        console_group = QGroupBox("Activity Log")
        console_layout = QVBoxLayout(console_group)
        self.log_console = QTextEdit()
        self.log_console.setReadOnly(True)
        console_layout.addWidget(self.log_console)
        splitter.addWidget(console_group)

        main_layout.addWidget(splitter)
        splitter.setSizes([550, 300])

        self._reset_context_combo()

        # --- Connect Signals ---
        self.btn_browse.clicked.connect(self._browse_chromedriver)
        self.btn_download_driver.clicked.connect(self._action_download_chromedriver)
        self.btn_open.clicked.connect(self._launch_or_navigate)
        self.btn_close.clicked.connect(self._close_browser)
        self.btn_scan_contexts.clicked.connect(self._action_scan_contexts)
        self.btn_find.clicked.connect(self._action_find_elements)
        self.results_combo.activated.connect(self._on_element_selected)
        self.btn_click.clicked.connect(self._action_click)
        self.btn_get_text.clicked.connect(self._action_get_text)
        self.btn_send_keys.clicked.connect(self._action_send_keys)

        self.btn_save_config.clicked.connect(self._action_save_config)
        self.btn_load_config.clicked.connect(self._action_load_config_dialog)
        self.btn_reset_config.clicked.connect(self._action_reset_config)
        self.btn_take_image.clicked.connect(self._action_launch_take_image)
        self.btn_manage_web_list.clicked.connect(self._action_open_web_list_editor)

        # Load default or saved configuration on startup
        self._load_default_or_saved_config()

    def _action_open_web_list_editor(self):
        """Opens the QDialog modal to view, edit, add, delete and save list_web.xlsx configuration."""
        dialog = WebListEditorDialog(self, excel_path="list_web.xlsx")
        dialog.exec()

    # --- Get Price Action Handlers ---
    def _action_start_get_price_workflow(self):
        if self.get_price_thread and self.get_price_thread.isRunning():
            return

        product_name = self.get_price_product_input.text().strip()
        if not product_name:
            return self.log("Product Name cannot be empty.", "WARN")

        cfg = self.get_config_dict()
        process_first_only = self.chk_process_first_only.isChecked()

        self.btn_start_get_price.setEnabled(False)
        self.btn_start_get_price.setText("⏳ Running...")

        self.lbl_extracted_name.setText("Processing...")
        self.lbl_extracted_price.setText("Processing...")
        self.lbl_extracted_status.setText("Processing...")
        self.lbl_extracted_similarity.setText("Processing...")
        self.lbl_extracted_link.clear()
        self.txt_raw_gemini_output.clear()
        self.lbl_screenshot_preview.setText("Capturing screenshot...")
        self.combo_processed_links.clear()
        self.lbl_price_summary.setText("Cheapest: - | Most Expensive: -")

        mode_text = "1st priority link only" if process_first_only else "ALL matching priority links"
        self.log(f"Starting automated 'Get Price' workflow for: '{product_name}' ({mode_text})", "HEADING")

        self.get_price_thread = GetPriceWorkflowThread(product_name, cfg, process_first_only)
        self.get_price_thread.log_signal.connect(self.log)
        self.get_price_thread.progress_signal.connect(self._on_get_price_progress)
        self.get_price_thread.search_table_signal.connect(self._on_update_search_table_from_workflow)
        self.get_price_thread.finished_signal.connect(self._on_get_price_finished)
        self.get_price_thread.start()

    def _on_get_price_progress(self, status_text: str):
        self.get_price_status_label.setText(status_text)

    def _on_update_search_table_from_workflow(self, items: list):
        self._on_google_search_finished(True, items, "")

    def _on_get_price_finished(self, success: bool, results_list: list, err_msg: str, final_prompt: str):
        self.btn_start_get_price.setEnabled(True)
        self.btn_start_get_price.setText("🚀 Get Price")

        self.get_price_results = results_list
        self.combo_processed_links.blockSignals(True)
        self.combo_processed_links.clear()

        if not success or not results_list:
            self.lbl_extracted_name.setText("Failed / Not Found")
            self.lbl_extracted_price.setText("Not Found")
            self.lbl_extracted_status.setText("Stopped")
            self.lbl_extracted_similarity.setText("0%")
            self.lbl_extracted_link.clear()
            self.txt_raw_gemini_output.setPlainText(err_msg)
            self.lbl_price_summary.setText("Cheapest: - | Most Expensive: -")
            self.combo_processed_links.blockSignals(False)
            self.log(f"Get Price workflow stopped: {err_msg}", "ERROR")
            return

        # Calculate cheapest and most expensive
        valid_items = [r for r in results_list if r.get("numeric_price", float('inf')) != float('inf')]
        self.cheapest_idx = -1
        self.expensive_idx = -1

        if valid_items:
            min_item = min(valid_items, key=lambda x: x["numeric_price"])
            max_item = max(valid_items, key=lambda x: x["numeric_price"])
            self.cheapest_idx = results_list.index(min_item)
            self.expensive_idx = results_list.index(max_item)

            self.lbl_price_summary.setText(
                f"<b>🏷️ Cheapest:</b> <span style='color: green;'>{min_item['price']}</span> ({min_item['domain']}) &nbsp;|&nbsp; "
                f"<b>💎 Most Expensive:</b> <span style='color: #d9534f;'>{max_item['price']}</span> ({max_item['domain']})"
            )
        else:
            self.lbl_price_summary.setText("Cheapest: N/A | Most Expensive: N/A")

        for idx, item in enumerate(results_list):
            badge = ""
            if idx == self.cheapest_idx and self.cheapest_idx != -1:
                badge += " 🏷️ [CHEAPEST]"
            if idx == self.expensive_idx and self.expensive_idx != -1:
                badge += " 💎 [MOST EXPENSIVE]"

            disp = f"[{idx+1}/{len(results_list)}] {item['domain']} - {item['price']}{badge}"
            self.combo_processed_links.addItem(disp, idx)

        self.combo_processed_links.blockSignals(False)

        # Select cheapest link by default if available, else first link
        default_idx = self.cheapest_idx if self.cheapest_idx >= 0 else 0
        self.combo_processed_links.setCurrentIndex(default_idx)
        self._on_processed_link_changed(default_idx)

        self.log(f"Automated 'Get Price' workflow completed! Processed {len(results_list)} link(s).", "SUCCESS")

    def _on_processed_link_changed(self, index: int):
        if index < 0 or index >= len(self.get_price_results):
            return

        item = self.get_price_results[index]
        self.lbl_extracted_name.setText(str(item.get("product_name", "-")))
        self.lbl_extracted_price.setText(str(item.get("price", "-")))
        self.lbl_extracted_status.setText(str(item.get("status", "-")))
        self.lbl_extracted_similarity.setText(str(item.get("similarity", "-")))
        self.lbl_extracted_link.setText(str(item.get("link", "-")))
        self.txt_raw_gemini_output.setPlainText(str(item.get("raw_text", "")))

        screenshot_path = item.get("screenshot_path", "")
        if screenshot_path and os.path.exists(screenshot_path):
            pixmap = QPixmap(screenshot_path)
            if not pixmap.isNull():
                self.lbl_screenshot_preview.setPixmap(
                    pixmap.scaled(480, 420, Qt.AspectRatioMode.KeepAspectRatio, Qt.TransformationMode.SmoothTransformation)
                )
            else:
                self.lbl_screenshot_preview.setText(f"Saved: {os.path.basename(screenshot_path)}")

        # Update Tab 5 (Google Gemini Test) so user can inspect
        final_prompt = item.get("final_prompt", "")
        if final_prompt:
            self.gemini_prompt_input.setPlainText(final_prompt)
        if screenshot_path:
            self.gemini_media_input.setText(screenshot_path)
        if item.get("raw_text"):
            self.gemini_output_text.setPlainText(item.get("raw_text", ""))

    def _select_cheapest_link(self):
        if self.cheapest_idx >= 0 and self.cheapest_idx < self.combo_processed_links.count():
            self.combo_processed_links.setCurrentIndex(self.cheapest_idx)

    def _select_expensive_link(self):
        if self.expensive_idx >= 0 and self.expensive_idx < self.combo_processed_links.count():
            self.combo_processed_links.setCurrentIndex(self.expensive_idx)

    def _open_screenshot_folder(self):
        folder = os.path.abspath("screenshot")
        os.makedirs(folder, exist_ok=True)
        if platform.system() == "Windows":
            os.startfile(folder)
        elif platform.system() == "Darwin":
            subprocess.run(["open", folder])
        else:
            subprocess.run(["xdg-open", folder])

    # --- Google Search Action Handlers ---
    def _action_run_google_search(self):
        if self.search_thread and self.search_thread.isRunning():
            return

        api_key = self.google_search_api_key
        cx = self.google_search_engine_id
        query = self.search_query_input.text().strip()

        if not api_key:
            return self.log("Google Search API Key is empty. Configure it in the Configuration tab.", "WARN")
        if not cx:
            return self.log("Custom Search Engine ID (CX) is empty. Configure it in the Configuration tab.", "WARN")
        if not query:
            return self.log("Search query cannot be empty.", "WARN")

        self.btn_run_search.setEnabled(False)
        self.btn_run_search.setText("⏳ Searching...")
        self.search_results_table.setRowCount(0)
        self.log(f"Initiating Google Search query: '{query}'", "HEADING")

        self.search_thread = GoogleSearchThread(api_key, cx, query)
        self.search_thread.log_signal.connect(self.log)
        self.search_thread.results_signal.connect(self._on_google_search_finished)
        self.search_thread.start()

    def _on_google_search_finished(self, success: bool, items: list, err_msg: str):
        self.btn_run_search.setEnabled(True)
        self.btn_run_search.setText("🔍 Search Google")

        if not success:
            self.log(f"Search failed: {err_msg}", "ERROR")
            return

        self.search_results_table.setRowCount(len(items))
        for row_idx, item in enumerate(items):
            title_item = QTableWidgetItem(item.get("title", ""))
            snippet_item = QTableWidgetItem(item.get("snippet", ""))
            link_item = QTableWidgetItem(item.get("link", ""))

            self.search_results_table.setItem(row_idx, 0, title_item)
            self.search_results_table.setItem(row_idx, 1, snippet_item)
            self.search_results_table.setItem(row_idx, 2, link_item)

        self.log(f"Google Search returned {len(items)} result(s). Data table updated.", "SUCCESS")

    # --- Google Gemini Action Handlers ---
    def _browse_gemini_media_file(self):
        fname, _ = QFileDialog.getOpenFileName(
            self,
            "Select Image or PDF Document",
            "",
            "Supported Files (*.png *.jpg *.jpeg *.webp *.pdf);;Images (*.png *.jpg *.jpeg *.webp);;PDF Documents (*.pdf);;All Files (*)"
        )
        if fname:
            self.gemini_media_input.setText(fname)

    def _action_send_gemini_prompt(self):
        if self.gemini_thread and self.gemini_thread.isRunning():
            return

        api_key = self.gemini_api_key
        model = self.gemini_model
        prompt = self.gemini_prompt_input.toPlainText().strip()
        media_path = self.gemini_media_input.text().strip()

        if not api_key:
            return self.log("Gemini API Key is empty. Configure it in the Configuration tab.", "WARN")
        if not prompt and not media_path:
            return self.log("Please enter a prompt or attach a file.", "WARN")

        self.btn_send_gemini.setEnabled(False)
        self.btn_send_gemini.setText("⏳ Processing...")
        self.gemini_output_text.clear()
        self.log(f"Sending prompt to Gemini ({model})...", "HEADING")

        self.gemini_thread = GeminiApiThread(api_key, model, prompt, media_path)
        self.gemini_thread.log_signal.connect(self.log)
        self.gemini_thread.response_signal.connect(self._on_gemini_finished)
        self.gemini_thread.start()

    def _on_gemini_finished(self, success: bool, output_text: str):
        self.btn_send_gemini.setEnabled(True)
        self.btn_send_gemini.setText("🚀 Send Prompt to Gemini")

        if success:
            self.gemini_output_text.setPlainText(output_text)
            self.log("Gemini response received successfully!", "SUCCESS")
        else:
            self.gemini_output_text.setPlainText(f"ERROR: {output_text}")
            self.log(f"Gemini API request failed: {output_text}", "ERROR")

    def _browse_working_folder(self):
        curr = self.working_folder_input.text().strip() or os.getcwd()
        folder = QFileDialog.getExistingDirectory(self, "Select Working Folder", curr)
        if folder:
            self.working_folder_input.setText(folder)
            self.log(f"Working folder updated to: {folder}", "INFO")

    def _toggle_gemini_key_visibility(self, checked: bool):
        if checked:
            self.gemini_key_input.setEchoMode(QLineEdit.EchoMode.Normal)
            self.btn_toggle_gemini_key.setText("🙈 Hide")
        else:
            self.gemini_key_input.setEchoMode(QLineEdit.EchoMode.Password)
            self.btn_toggle_gemini_key.setText("👁️ Show")

    def get_config_dict(self) -> Dict[str, Any]:
        return {
            "chromedriver_path": self.driver_path_input.text().strip(),
            "working_folder": self.working_folder_input.text().strip(),
            "gemini_api_key": self.gemini_key_input.text().strip(),
            "gemini_model": self.gemini_model_combo.currentText().strip(),
            "gemini_delay_between_calls": float(self.spin_gemini_delay_between_calls.value()),
            "gemini_retry_delay": float(self.spin_gemini_retry_delay.value()),
            "gemini_max_retries": int(self.spin_gemini_max_retries.value()),
            "google_search_api_key": self.search_key_input.text().strip(),
            "google_search_engine_id": self.search_cx_input.text().strip(),
            "target_url": self.url_input.text().strip(),
        }

    def apply_config_dict(self, cfg: Dict[str, Any]):
        driver_path = cfg.get("chromedriver_path", "")
        if not driver_path:
            default_driver = os.path.abspath("chromedriver.exe" if platform.system() == "Windows" else "chromedriver")
            if os.path.exists(default_driver):
                driver_path = default_driver

        self.driver_path_input.setText(driver_path)
        self.working_folder_input.setText(cfg.get("working_folder", os.getcwd()))
        self.gemini_key_input.setText(cfg.get("gemini_api_key", ""))

        model_name = cfg.get("gemini_model", "gemini-2.5-flash")
        idx = self.gemini_model_combo.findText(model_name)
        if idx >= 0:
            self.gemini_model_combo.setCurrentIndex(idx)
        else:
            self.gemini_model_combo.setEditText(model_name)

        self.spin_gemini_delay_between_calls.setValue(float(cfg.get("gemini_delay_between_calls", 2.0)))
        self.spin_gemini_retry_delay.setValue(float(cfg.get("gemini_retry_delay", 5.0)))
        self.spin_gemini_max_retries.setValue(int(cfg.get("gemini_max_retries", 3)))

        self.search_key_input.setText(cfg.get("google_search_api_key", ""))
        self.search_cx_input.setText(cfg.get("google_search_engine_id", ""))
        self.url_input.setText(cfg.get("target_url", "https://www.google.com"))

    def _load_default_or_saved_config(self):
        if os.path.exists(CONFIG_FILE_PATH):
            try:
                with open(CONFIG_FILE_PATH, "r", encoding="utf-8") as f:
                    cfg = json.load(f)
                self.apply_config_dict(cfg)
                self.log(f"Loaded configuration from: {CONFIG_FILE_PATH}", "SUCCESS")
                return
            except Exception as e:
                self.log(f"Failed to read config file ({CONFIG_FILE_PATH}): {e}", "WARN")

        self.apply_config_dict(DEFAULT_CONFIG)
        self.log("Loaded default configuration.", "INFO")

    def _action_save_config(self):
        try:
            cfg = self.get_config_dict()
            with open(CONFIG_FILE_PATH, "w", encoding="utf-8") as f:
                json.dump(cfg, f, indent=4)
            self.log(f"Configuration saved to: {CONFIG_FILE_PATH}", "SUCCESS")
        except Exception as e:
            self.log(f"Failed to save configuration: {e}", "ERROR")

    def _action_load_config_dialog(self):
        fname, _ = QFileDialog.getOpenFileName(self, "Select Config JSON File", "", "JSON Files (*.json)")
        if fname:
            try:
                with open(fname, "r", encoding="utf-8") as f:
                    cfg = json.load(f)
                self.apply_config_dict(cfg)
                self.log(f"Loaded configuration from: {fname}", "SUCCESS")
            except Exception as e:
                self.log(f"Failed to load config file: {e}", "ERROR")

    def _action_reset_config(self):
        self.apply_config_dict(DEFAULT_CONFIG)
        self._action_save_config()
        self.log("Reset configuration to default settings.", "SUCCESS")

    def _action_launch_take_image(self):
        bot_script = os.path.join(os.path.dirname(os.path.abspath(__file__)), "BOT_take_image.py")
        if not os.path.exists(bot_script):
            QMessageBox.critical(self, "Error", f"Script not found: {bot_script}")
            return

        self.log("Minimizing main app and launching BOT_take_image...", "INFO")
        self.showMinimized()

        self.bot_process = QProcess(self)
        self.bot_process.finished.connect(self._on_bot_take_image_finished)
        self.bot_process.start(sys.executable, [bot_script])

    def _on_bot_take_image_finished(self, exit_code, exit_status):
        self.log(f"BOT_take_image process finished (exit code: {exit_code}). Restoring main app...", "INFO")
        if exit_code != 0 and hasattr(self, 'bot_process'):
            err = self.bot_process.readAllStandardError().data().decode("utf-8", errors="replace")
            if err.strip():
                self.log(f"BOT_take_image error output: {err.strip()}", "ERROR")
                QMessageBox.warning(self, "Tool Error", f"BOT_take_image closed with error code {exit_code}:\n\n{err.strip()}")
        self.showNormal()
        self.activateWindow()
        self.raise_()

    @property
    def working_folder(self) -> str:
        return self.working_folder_input.text().strip() or os.getcwd()

    @property
    def chromedriver_path(self) -> str:
        return self.driver_path_input.text().strip()

    @property
    def gemini_api_key(self) -> str:
        return self.gemini_key_input.text().strip()

    @property
    def gemini_model(self) -> str:
        return self.gemini_model_combo.currentText().strip()

    @property
    def google_search_api_key(self) -> str:
        return self.search_key_input.text().strip()

    @property
    def google_search_engine_id(self) -> str:
        return self.search_cx_input.text().strip()

    def _reset_context_combo(self):
        self.context_combo.clear()
        self.context_cache.clear()
        self.context_combo.addItem("Default Page Content", -1)

    def log(self, text: str, level: str = "INFO"):
        stamp = time.strftime("%H:%M:%S")
        color = {"ERROR": "red", "SUCCESS": "green", "WARN": "orange", "HEADING": "blue"}.get(level, "black")
        self.log_console.append(f'<span style="color: gray;">[{stamp}]</span> <b style="color: {color};">[{level}]</b> {text}')

    def _browse_chromedriver(self):
        fname, _ = QFileDialog.getOpenFileName(self, "Select chromedriver", "", "Executables (*.exe *)")
        if fname: self.driver_path_input.setText(fname)

    def _action_download_chromedriver(self):
        if self.download_thread and self.download_thread.isRunning():
            return
        
        self.btn_download_driver.setEnabled(False)
        self.btn_download_driver.setText("⏳ Downloading...")
        self.log("Starting automatic ChromeDriver downloader...", "HEADING")

        self.download_thread = ChromeDriverDownloaderThread()
        self.download_thread.log_signal.connect(self.log)
        self.download_thread.finished_signal.connect(self._on_download_finished)
        self.download_thread.start()

    def _on_download_finished(self, success: bool, result: str):
        self.btn_download_driver.setEnabled(True)
        self.btn_download_driver.setText("⬇️ Auto-Download ChromeDriver")
        if success:
            self.driver_path_input.setText(result)
            self.log(f"ChromeDriver downloaded successfully: {result}", "SUCCESS")
            self.log("ChromeDriver path updated automatically. You can now open browser!", "INFO")
        else:
            self.log(f"ChromeDriver download failed: {result}", "ERROR")

    def _launch_or_navigate(self):
        url = self.url_input.text().strip()
        if not url: return self.log("URL cannot be empty.", "WARN")
        try:
            if not self.driver:
                path = self.driver_path_input.text().strip()
                service = ChromeService(executable_path=path) if path else ChromeService()
                self.driver = webdriver.Chrome(service=service)
            self.driver.get(url)
            self.log(f"Loaded: {self.driver.title}", "SUCCESS")
            self._reset_context_combo()
        except WebDriverException as e:
            self.log(f"Launch/Navigate failed: {e.msg}", "ERROR")

    def _close_browser(self):
        self._clear_highlight()
        if self.driver:
            try: self.driver.quit()
            except Exception: pass
            self.driver = None
        self.found_elements_cache = []
        self.results_combo.clear()
        self.selected_locator_input.clear()
        self._reset_context_combo()
        self.log("Browser closed.", "INFO")

    def _action_scan_contexts(self):
        if not self.driver: return self.log("Browser not open.", "ERROR")
        self.log("Scanning page for iframes and Shadow DOMs...", "INFO")
        try:
            self.driver.switch_to.default_content()
            results = self.driver.execute_script(JS_LIST_CONTEXTS)
            self._reset_context_combo()
            
            for item in results.get('iframes', []):
                item['type'] = 'iframe'
                self.context_cache.append(item)
                display_text = f"IFRAME: #{item['id'] or item['name'] or 'unknown'}"
                self.context_combo.addItem(display_text, len(self.context_cache) - 1)

            for item in results.get('shadow_hosts', []):
                item['type'] = 'shadow_host'
                self.context_cache.append(item)
                display_text = f"SHADOW: <{item['host_tag']}> #{item['host_id'] or 'unknown'}"
                self.context_combo.addItem(display_text, len(self.context_cache) - 1)
            
            self.log(f"Scan complete. Found {len(self.context_cache)} contexts.", "SUCCESS")
        except Exception as e: self.log(f"Failed to scan page: {e}", "ERROR")

    def _action_find_elements(self):
        if not self.driver: return self.log("Browser not open.", "ERROR")
        
        strategy = self.combo_by.currentText()
        selector = self.selector_input.text().strip()
        if not selector: return self.log("Selector/Text cannot be empty.", "WARN")

        self.results_combo.clear(); self.found_elements_cache = []
        self._clear_highlight()

        context_idx = self.context_combo.currentData()
        context_info = self.context_cache[context_idx] if context_idx != -1 else {"type": "default"}
        
        self.log(f"Finding elements by {strategy} in context: {self.context_combo.currentText()}", "HEADING")
        
        try:
            self.driver.switch_to.default_content()
            elements = []
            search_args = [selector, 'smart' if strategy == "Smart Search (Text)" else 'contains']

            if context_info['type'] == 'iframe':
                frame_element = self.driver.find_element(By.XPATH, context_info['xpath'])
                self.driver.switch_to.frame(frame_element)
                if strategy == "Smart Search (Text)":
                    elements = self.driver.execute_script(JS_SMART_SEARCH_ENGINE, *search_args)
                else:
                    elements = self.driver.find_elements(getattr(By, strategy), selector)

            elif context_info['type'] == 'shadow_host':
                host_element = self.driver.find_element(By.XPATH, context_info['host_xpath'])
                search_args.append(host_element.shadow_root)
                if strategy == "Smart Search (Text)":
                    elements = self.driver.execute_script(JS_SMART_SEARCH_ENGINE, *search_args)
                else: # Fallback for non-JS search in shadow root
                    elements = host_element.shadow_root.find_elements(getattr(By, strategy.replace("_", " ")), selector)

            else: # Default content
                if strategy == "Smart Search (Text)":
                    elements = self.driver.execute_script(JS_SMART_SEARCH_ENGINE, *search_args)
                else:
                    elements = self.driver.find_elements(getattr(By, strategy), selector)

            if not elements: return self.log("No elements found in this context.", "WARN")

            self.found_elements_cache = elements
            self.log(f"Found {len(elements)} element(s). Populating dropdown.", "SUCCESS")
            for i, el in enumerate(elements):
                try:
                    text = (el.text or el.get_attribute('value') or f"<{el.tag_name}>").strip()[:60]
                    self.results_combo.addItem(f"{i}: {text}", i)
                except Exception: self.results_combo.addItem(f"{i}: [Stale Element]", i)
            self.results_combo.setCurrentIndex(0)
            self._on_element_selected()

        except Exception as e: self.log(f"Find failed: {e}", "ERROR")
        finally:
            if self.driver: self.driver.switch_to.default_content()

    def _switch_to_correct_context(self) -> None:
        """
        Switches the driver to the context selected in the dropdown before performing an action.
        """
        if not self.driver:
            return
            
        self.driver.switch_to.default_content() # Start from a clean slate
        context_idx = self.context_combo.currentData()
        if context_idx is None or context_idx == -1:
            return # In default content, do nothing

        context_info = self.context_cache[context_idx]
        if context_info['type'] == 'iframe':
            try:
                frame_element = self.driver.find_element(By.XPATH, context_info['xpath'])
                self.driver.switch_to.frame(frame_element)
                self.log(f"Switched to context: {self.context_combo.currentText()}", "INFO")
            except Exception as e:
                self.log(f"Could not switch back to iframe context: {e}", "ERROR")

    def _on_element_selected(self):
        if not self.driver or not self.found_elements_cache: return
        index = self.results_combo.currentData()
        if index is None: return

        element = self.found_elements_cache[index]
        self._clear_highlight() 

        try:
            self._switch_to_correct_context()
            
            self.driver.execute_script(JS_HIGHLIGHT_ELEMENT, element, 'blue')
            self.highlighted_element = element
            
            is_shadow = self.driver.execute_script("return arguments[0].getRootNode() instanceof ShadowRoot", element)
            prefix = "(in Shadow DOM) > " if is_shadow else ""
            xpath = self.driver.execute_script(JS_GET_XPATH, element)
            self.selected_locator_input.setText(prefix + xpath)

        except Exception as e:
            self.log(f"Failed to inspect element: {e}", "WARN")
        finally:
            if self.driver:
                self.driver.switch_to.default_content()

    def _get_selected_element(self) -> Optional[WebElement]:
        if not self.highlighted_element:
            self.log("No element selected.", "WARN")
            return None
        return self.highlighted_element

    def _action_click(self):
        el = self._get_selected_element()
        if el:
            try:
                self._switch_to_correct_context()
                self.log("Attempting native Selenium click...", "INFO")
                el.click()
                self.log("Native click was successful.", "SUCCESS")
                
            except Exception as e:
                self.log(f"Native click failed: {type(e).__name__}. Trying JavaScript click as a fallback.", "WARN")
                try:
                    self.driver.execute_script("arguments[0].click();", el)
                    self.log("JavaScript fallback click was successful.", "SUCCESS")
                except Exception as e2:
                    self.log(f"JavaScript fallback click also failed: {e2}", "ERROR")
            finally:
                if self.driver:
                    self.driver.switch_to.default_content()

    def _action_get_text(self):
        el = self._get_selected_element()
        if el:
            try:
                # Switch to the context right before the action
                self._switch_to_correct_context()
                
                # Try getting the visible text first
                extracted_text = el.text
                
                # If text is empty (common for input/textarea fields), try retrieving the value
                if not extracted_text:
                    extracted_text = el.get_attribute('value') or el.get_attribute('textContent') or ""
                
                # Log the output to the activity log console
                self.log(f"Extracted Text: '{extracted_text.strip()}'", "SUCCESS")
                
            except Exception as e:
                self.log(f"Get Text failed: {e}", "ERROR")
            finally:
                # Always switch back after the action is complete
                if self.driver:
                    self.driver.switch_to.default_content()

    def _action_send_keys(self):
        el = self._get_selected_element()
        text_to_send = self.value_input.text()
        if el and text_to_send:
            try:
                self._switch_to_correct_context()
                el.send_keys(text_to_send)
                self.log(f"Sent keys: '{text_to_send}'", "SUCCESS")
            except Exception as e:
                self.log(f"Send Keys failed: {e}", "ERROR")
            finally:
                if self.driver:
                    self.driver.switch_to.default_content()

    def _clear_highlight(self):
        if self.highlighted_element and self.driver:
            try: 
                self.driver.switch_to.default_content()
                self.driver.execute_script(JS_CLEAR_HIGHLIGHT, self.highlighted_element)
            except Exception: 
                pass 
        self.highlighted_element = None
        
    # --- Tab 0: Get Price All (Batch Mode) Methods ---
    def _init_tab_get_price_all(self) -> QWidget:
        tab_get_price_all = QWidget()
        layout = QVBoxLayout(tab_get_price_all)

        top_layout = QHBoxLayout()

        # Top Left: Excel File Loader
        group_left = QGroupBox("1. Excel Data Source Setup")
        left_layout = QGridLayout(group_left)

        left_layout.addWidget(QLabel("Excel File Path:"), 0, 0)
        self.txt_excel_path_all = QLineEdit()
        self.txt_excel_path_all.setReadOnly(True)
        self.txt_excel_path_all.setPlaceholderText("Click button to load Excel file (e.g. Get_price_list.xlsx)...")
        left_layout.addWidget(self.txt_excel_path_all, 0, 1)

        self.btn_load_excel_all = QPushButton("📂 Load Excel File")
        self.btn_load_excel_all.setStyleSheet("font-weight: bold;")
        self.btn_load_excel_all.clicked.connect(self._action_load_excel_all)
        left_layout.addWidget(self.btn_load_excel_all, 0, 2)

        self.lbl_excel_info_all = QLabel("No file loaded. Required template: Get_price_list.xlsx schema.")
        self.lbl_excel_info_all.setStyleSheet("font-size: 11px; color: gray;")
        left_layout.addWidget(self.lbl_excel_info_all, 1, 0, 1, 3)

        top_layout.addWidget(group_left, stretch=5)

        # Top Right: Execution Range & Control Panel
        group_right = QGroupBox("2. Execution Range & Control Panel")
        right_layout = QGridLayout(group_right)

        self.rdo_range_all = QRadioButton("Get price of whole file (Skip 'done' rows)")
        self.rdo_range_all.setChecked(True)
        self.rdo_range_all.toggled.connect(self._on_range_option_changed)
        right_layout.addWidget(self.rdo_range_all, 0, 0, 1, 4)

        self.rdo_range_custom = QRadioButton("From line:")
        self.rdo_range_custom.toggled.connect(self._on_range_option_changed)
        right_layout.addWidget(self.rdo_range_custom, 1, 0)

        self.spn_start_line = QSpinBox()
        self.spn_start_line.setRange(1, 9999)
        self.spn_start_line.setValue(1)
        self.spn_start_line.setEnabled(False)
        right_layout.addWidget(self.spn_start_line, 1, 1)

        right_layout.addWidget(QLabel("to line:"), 1, 2)

        self.spn_end_line = QSpinBox()
        self.spn_end_line.setRange(1, 9999)
        self.spn_end_line.setValue(1)
        self.spn_end_line.setEnabled(False)
        right_layout.addWidget(self.spn_end_line, 1, 3)

        # Action Buttons Layout
        btn_bar = QHBoxLayout()
        self.btn_start_batch = QPushButton("▶️ Start")
        self.btn_start_batch.setStyleSheet("font-weight: bold; background-color: #2b5797; color: white; padding: 4px 12px;")
        self.btn_start_batch.clicked.connect(self._action_start_batch)
        btn_bar.addWidget(self.btn_start_batch)

        self.btn_pause_batch = QPushButton("⏸️ Pause")
        self.btn_pause_batch.setEnabled(False)
        self.btn_pause_batch.clicked.connect(self._action_toggle_pause_batch)
        btn_bar.addWidget(self.btn_pause_batch)

        self.btn_stop_batch = QPushButton("⏹️ Stop")
        self.btn_stop_batch.setEnabled(False)
        self.btn_stop_batch.clicked.connect(self._action_stop_batch)
        btn_bar.addWidget(self.btn_stop_batch)

        right_layout.addLayout(btn_bar, 2, 0, 1, 4)

        top_layout.addWidget(group_right, stretch=5)
        layout.addLayout(top_layout)

        # Center Section: Data Table
        group_table = QGroupBox("3. Product Price Batch Data Table")
        table_layout = QVBoxLayout(group_table)

        self.tbl_batch_data = QTableWidget()
        self.tbl_batch_data.setColumnCount(len(REQUIRED_EXCEL_COLUMNS))
        self.tbl_batch_data.setHorizontalHeaderLabels(REQUIRED_EXCEL_COLUMNS)
        self.tbl_batch_data.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        self.tbl_batch_data.horizontalHeader().setStretchLastSection(True)
        table_layout.addWidget(self.tbl_batch_data)

        layout.addWidget(group_table)

        # Bottom Status Bar
        self.lbl_batch_status = QLabel("Ready. Please click '📂 Load Excel File' to select your data template.")
        self.lbl_batch_status.setStyleSheet("font-weight: bold; color: #2b5797;")
        layout.addWidget(self.lbl_batch_status)

        return tab_get_price_all

    def _action_load_excel_all(self):
        default_dir = os.getcwd()
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "Select Get Price List Excel File",
            default_dir,
            "Excel Files (*.xlsx *.xls)"
        )

        if not file_path:
            return

        is_valid, err_msg, headers, rows_data = validate_and_read_excel(file_path)

        if not is_valid:
            QMessageBox.critical(self, "Invalid Excel Template", err_msg)
            return

        self.batch_excel_path = file_path
        self.batch_rows_data = rows_data
        self.batch_headers = headers

        self.txt_excel_path_all.setText(file_path)
        self.lbl_excel_info_all.setText(f"Loaded {len(rows_data)} row(s) from {os.path.basename(file_path)}")
        self.lbl_batch_status.setText(f"Excel file loaded successfully with {len(rows_data)} item(s).")

        total_count = max(1, len(rows_data))
        self.spn_start_line.setRange(1, total_count)
        self.spn_end_line.setRange(1, total_count)
        self.spn_start_line.setValue(1)
        self.spn_end_line.setValue(total_count)

        self._populate_batch_table()

    def _populate_batch_table(self):
        from PyQt6.QtGui import QColor

        self.tbl_batch_data.setRowCount(0)
        self.tbl_batch_data.setRowCount(len(self.batch_rows_data))

        for row_idx, item_info in enumerate(self.batch_rows_data):
            row_dict = item_info["data"]
            status_val = str(row_dict.get("status", "")).strip().lower()

            for col_idx, col_name in enumerate(REQUIRED_EXCEL_COLUMNS):
                val = str(row_dict.get(col_name, ""))
                item = QTableWidgetItem(val)
                if status_val == "done":
                    item.setBackground(QColor("#e8f5e9"))
                self.tbl_batch_data.setItem(row_idx, col_idx, item)

    def _on_range_option_changed(self):
        is_custom = self.rdo_range_custom.isChecked()
        self.spn_start_line.setEnabled(is_custom)
        self.spn_end_line.setEnabled(is_custom)

    def _action_start_batch(self):
        if not self.batch_excel_path or not self.batch_rows_data:
            QMessageBox.warning(self, "No Data Loaded", "Please load a valid Excel file before starting batch execution.")
            return

        if self.batch_thread and self.batch_thread.isRunning():
            return

        if self.rdo_range_all.isChecked():
            eligible_rows = [item for item in self.batch_rows_data if str(item["data"].get("status", "")).strip().lower() != "done"]
        else:
            start_l = self.spn_start_line.value()
            end_l = self.spn_end_line.value()
            if start_l > end_l:
                QMessageBox.warning(self, "Invalid Range", "Start line cannot be greater than End line.")
                return
            eligible_rows = []
            for item in self.batch_rows_data:
                line_no = item["ui_index"] + 1
                if start_l <= line_no <= end_l:
                    if str(item["data"].get("status", "")).strip().lower() != "done":
                        eligible_rows.append(item)

        if not eligible_rows:
            QMessageBox.information(
                self,
                "No Rows to Process",
                "No pending rows found in selected range. All rows are already marked as 'done' or outside range."
            )
            return

        config = self.get_config_dict()
        self.batch_thread = GetPriceAllWorkerThread(self.batch_excel_path, eligible_rows, config)
        self.batch_thread.log_signal.connect(self.log)
        self.batch_thread.progress_signal.connect(self._on_batch_progress)
        self.batch_thread.row_started_signal.connect(self._on_batch_row_started)
        self.batch_thread.row_finished_signal.connect(self._on_batch_row_finished)
        self.batch_thread.finished_signal.connect(self._on_batch_finished)

        self.btn_start_batch.setEnabled(False)
        self.btn_pause_batch.setEnabled(True)
        self.btn_pause_batch.setText("⏸️ Pause")
        self.btn_stop_batch.setEnabled(True)
        self.btn_load_excel_all.setEnabled(False)

        self.batch_thread.start()

    def _on_batch_progress(self, msg: str):
        self.lbl_batch_status.setText(msg)

    def _action_toggle_pause_batch(self):
        if not self.batch_thread or not self.batch_thread.isRunning():
            return

        if self.batch_thread._is_paused:
            self.batch_thread.resume()
            self.btn_pause_batch.setText("⏸️ Pause")
            self.lbl_batch_status.setText("Resumed batch processing...")
            self.log("Batch execution resumed.", "INFO")
        else:
            self.batch_thread.pause()
            self.btn_pause_batch.setText("▶️ Resume")
            self.lbl_batch_status.setText("Batch processing paused.")
            self.log("Batch execution paused by user.", "WARN")

    def _action_stop_batch(self):
        if not self.batch_thread or not self.batch_thread.isRunning():
            return

        self.batch_thread.stop()
        self.btn_pause_batch.setEnabled(False)
        self.btn_stop_batch.setEnabled(False)
        self.lbl_batch_status.setText("Stopping batch execution...")
        self.log("Stop requested. Waiting for active step to complete...", "WARN")

    def _on_batch_row_started(self, ui_row_idx: int):
        from PyQt6.QtGui import QColor
        status_col = REQUIRED_EXCEL_COLUMNS.index("status")
        item = QTableWidgetItem("Processing...")
        item.setBackground(QColor("#fffde7"))
        self.tbl_batch_data.setItem(ui_row_idx, status_col, item)
        self.tbl_batch_data.scrollToItem(item)

    def _on_batch_row_finished(self, ui_row_idx: int, updated_row_dict: dict):
        from PyQt6.QtGui import QColor

        if ui_row_idx < len(self.batch_rows_data):
            self.batch_rows_data[ui_row_idx]["data"] = updated_row_dict

        for col_idx, col_name in enumerate(REQUIRED_EXCEL_COLUMNS):
            val = str(updated_row_dict.get(col_name, ""))
            item = QTableWidgetItem(val)
            item.setBackground(QColor("#e8f5e9"))
            self.tbl_batch_data.setItem(ui_row_idx, col_idx, item)

    def _on_batch_finished(self, success: bool, msg: str):
        self.btn_start_batch.setEnabled(True)
        self.btn_pause_batch.setEnabled(False)
        self.btn_pause_batch.setText("⏸️ Pause")
        self.btn_stop_batch.setEnabled(False)
        self.btn_load_excel_all.setEnabled(True)

        status_text = f"✅ {msg}" if success else f"⚠️ {msg}"
        self.lbl_batch_status.setText(status_text)
        self.log(msg, "SUCCESS" if success else "WARN")

    def closeEvent(self, event):
        self._close_browser()
        event.accept()

if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = ChromeDriverTesterApp()
    window.show()
    sys.exit(app.exec())